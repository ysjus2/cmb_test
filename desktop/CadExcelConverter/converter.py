from __future__ import annotations

import math
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import ezdxf
from ezdxf import recover
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pyproj import Transformer
from essenpoly_recovery import recover_essenpoly_polylines

LAYER_HEADERS = [
    "ENTITY_TYPE","ENTITY_ID","BLOCK_NAME","SEQ",
    "CAD_X","CAD_Y","경도","위도",
    "길이","TEXT","ATTRIBUTES","XDATA"
]

@dataclass
class LayerInfo:
    name: str
    count: int
    types: str

@dataclass
class ConversionStats:
    layers: int = 0
    entities: int = 0
    rows: int = 0
    skipped: int = 0

def _norm(value):
    return (value or "").strip()

def _validate_dxf(input_path):
    src = Path(input_path)
    if src.suffix.lower() != ".dxf":
        raise ValueError("회사용 최종판은 DXF 파일만 지원합니다.")
    if not src.is_file():
        raise FileNotFoundError(str(src))
    return src

def _sanitize_ascii_binary_tags(src, log):
    raw = Path(src).read_bytes().splitlines(keepends=True)
    out, fixed, i = [], 0, 0
    while i < len(raw):
        code_line = raw[i]
        if i + 1 >= len(raw):
            out.append(code_line)
            break
        value_line = raw[i + 1]
        try:
            code = int(code_line.strip().decode("ascii", errors="strict"))
        except Exception:
            out.append(code_line)
            i += 1
            continue
        if (310 <= code <= 319) or code == 1004:
            value = value_line.strip()
            valid = bool(value) and len(value) % 2 == 0 and all(
                ch in b"0123456789abcdefABCDEF" for ch in value
            )
            if not valid:
                newline = b"\r\n" if value_line.endswith(b"\r\n") else b"\n"
                out.extend([code_line, b"00" + newline])
                fixed += 1
                i += 2
                continue
        out.extend([code_line, value_line])
        i += 2

    temp = tempfile.NamedTemporaryFile(prefix="cmb_dxf_recover_", suffix=".dxf", delete=False)
    temp.write(b"".join(out))
    temp.close()
    log(f"DXF 복구용 비정상 바이너리 태그 {fixed}건 정리")
    return Path(temp.name)

def load_dxf_document(input_path, log: Callable[[str], None] | None = None):
    log = log or (lambda msg: None)
    src = _validate_dxf(input_path)

    # ezdxf 자체가 ASCII/Binary DXF를 모두 읽을 수 있으므로 정상 파일은 그대로 읽는다.
    try:
        log("DXF 읽기")
        return ezdxf.readfile(str(src))
    except Exception as first_error:
        log(f"일반 DXF 읽기 실패: {first_error}")

    # 손상된 ASCII DXF에 대해서만 복구본을 만든다.
    raw = src.read_bytes()
    if raw.startswith(b"AutoCAD Binary DXF"):
        raise RuntimeError("Binary DXF를 읽지 못했습니다. 정식 CAD에서 ASCII DXF로 다시 저장해주세요.")

    sanitized = None
    try:
        sanitized = _sanitize_ascii_binary_tags(src, log)
        try:
            doc, auditor = recover.readfile(str(sanitized), errors="ignore")
            log(f"DXF recover 읽기 완료 · 오류 {len(getattr(auditor, 'errors', []))}건")
            return doc
        except Exception as recover_error:
            log(f"recover 실패: {recover_error}")
        try:
            doc, auditor = recover.explore(str(sanitized), errors="ignore")
            log(f"DXF explore 읽기 완료 · 오류 {len(getattr(auditor, 'errors', []))}건")
            return doc
        except Exception as explore_error:
            raise RuntimeError(
                f"DXF를 읽지 못했습니다. recover={recover_error} / explore={explore_error}"
            ) from explore_error
    finally:
        if sanitized is not None:
            try:
                sanitized.unlink(missing_ok=True)
            except Exception:
                pass

def scan_layers(input_path, oda_exe=None, log=None, progress=None):
    log = log or (lambda msg: None)
    progress = progress or (lambda percent, task: None)
    progress(5, "DXF 읽기")
    doc = load_dxf_document(input_path, log)
    entities = list(doc.modelspace())
    counts, types = {}, {}
    total = max(1, len(entities))
    for index, ent in enumerate(entities, 1):
        layer = _norm(getattr(ent.dxf, "layer", "0")) or "0"
        counts[layer] = counts.get(layer, 0) + 1
        types.setdefault(layer, set()).add(ent.dxftype())
        if index == total or index % max(1, total // 100) == 0:
            progress(15 + int(index / total * 80), f"레이어 분석 {index:,}/{total:,}")
    result = [
        LayerInfo(name, counts[name], ", ".join(sorted(types[name])))
        for name in sorted(counts, key=str.lower)
    ]
    progress(100, f"레이어 {len(result)}개 분석 완료")
    return result

def _attributes(entity):
    out = {}
    if entity.dxftype() == "INSERT":
        for attr in getattr(entity, "attribs", []):
            out[_norm(attr.dxf.tag)] = _norm(attr.dxf.text)
    return out

def _joined_attrs(attrs):
    return " | ".join(f"{k}={v}" for k, v in attrs.items() if v)

def _xdata_text(doc, entity):
    chunks = []
    try:
        for appid in doc.appids:
            try:
                tags = entity.get_xdata(appid.dxf.name)
            except Exception:
                continue
            for tag in tags:
                if isinstance(tag.value, str) and tag.value.strip():
                    chunks.append(tag.value.strip())
    except Exception:
        pass
    return " | ".join(dict.fromkeys(chunks))

def _entity_text(entity):
    try:
        if entity.dxftype() == "TEXT":
            return _norm(entity.dxf.text)
        if entity.dxftype() == "MTEXT":
            return _norm(entity.plain_text())
    except Exception:
        pass
    return ""

def entity_points(entity):
    typ = entity.dxftype()
    try:
        if typ == "INSERT":
            p = entity.dxf.insert
            return [(float(p.x), float(p.y))]
        if typ == "POINT":
            p = entity.dxf.location
            return [(float(p.x), float(p.y))]
        if typ in {"TEXT", "MTEXT"}:
            p = entity.dxf.insert
            return [(float(p.x), float(p.y))]
        if typ == "LWPOLYLINE":
            return [(float(x), float(y)) for x, y, *_ in entity.get_points("xy")]
        if typ == "POLYLINE":
            return [(float(v.dxf.location.x), float(v.dxf.location.y)) for v in entity.vertices]
        if typ == "LINE":
            s, e = entity.dxf.start, entity.dxf.end
            return [(float(s.x), float(s.y)), (float(e.x), float(e.y))]
        if typ in {"CIRCLE", "ARC"}:
            p = entity.dxf.center
            return [(float(p.x), float(p.y))]
    except Exception:
        pass
    return []

def _length(points):
    return sum(
        math.hypot(b[0] - a[0], b[1] - a[1])
        for a, b in zip(points, points[1:])
    ) if len(points) >= 2 else 0.0

def _write_sheet(ws, headers, rows):
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, 1):
        sample = [len(str(header)) + 2]
        for r in range(2, min(ws.max_row, 300) + 1):
            sample.append(len(str(ws.cell(r, i).value or "")))
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(42, max(sample)))

def _safe_sheet_name(layer_name, used):
    base = re.sub(r'[\\/*?:\[\]]', "_", layer_name).strip() or "LAYER"
    base = base[:31]
    name, i = base, 2
    while name.lower() in used:
        suffix = f"_{i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name

def convert_selected_layers(
    input_path,
    output_path,
    selected_layers,
    source_epsg=5174,
    oda_exe=None,
    log=None,
    progress=None,
):
    log = log or (lambda msg: None)
    progress = progress or (lambda percent, task: None)
    selected = list(dict.fromkeys(selected_layers))
    if not selected:
        raise ValueError("추출할 레이어를 하나 이상 선택해주세요.")

    progress(5, "DXF 읽기")
    doc = load_dxf_document(input_path, log)
    transformer = Transformer.from_crs(f"EPSG:{source_epsg}", "EPSG:4326", always_xy=True)
    grouped = {name: [] for name in selected}
    entity_counts = {name: 0 for name in selected}
    stats = ConversionStats(layers=len(selected))

    # IMPORTANT:
    # 손상 DXF를 ezdxf recover로 읽으면 ESSENPOLY의 원래 레이어가 '0'으로
    # 바뀌는 도면이 있다. 따라서 ESSENPOLY는 modelspace 값을 신뢰하지 않고
    # 원본 ASCII DXF에서 복구한 실제 layer/handle/point/xdata를 직접 출력한다.
    recovered_items = recover_essenpoly_polylines(input_path)
    recovered_handles = {item.get("handle", "") for item in recovered_items if item.get("handle")}

    entities = list(doc.modelspace())
    total = max(1, len(entities))
    for index, ent in enumerate(entities, 1):
        typ = ent.dxftype()
        handle = _norm(getattr(ent.dxf, "handle", ""))

        # ESSENPOLY는 아래 원문 복구 단계에서만 처리한다.
        # 여기서 처리하면 recover가 만든 layer=0 빈 행이 Excel에 섞인다.
        if typ == "ESSENPOLY" or (handle and handle in recovered_handles):
            if index == total or index % max(1, total // 100) == 0:
                progress(15 + int(index / total * 55), f"객체 분석 {index:,}/{total:,}")
            continue

        layer = _norm(getattr(ent.dxf, "layer", "0")) or "0"
        if layer in grouped:
            entity_counts[layer] += 1
            stats.entities += 1
            block = _norm(getattr(ent.dxf, "name", "")) if typ == "INSERT" else ""
            attrs = _joined_attrs(_attributes(ent))
            xdata = _xdata_text(doc, ent)
            text_value = _entity_text(ent)
            points = entity_points(ent)
            length = _length(points)
            if points:
                for seq, (x, y) in enumerate(points, 1):
                    try:
                        lon, lat = transformer.transform(x, y)
                    except Exception:
                        lon, lat = None, None
                    grouped[layer].append([
                        typ, handle, block, seq, x, y, lon, lat,
                        length, text_value, attrs, xdata,
                    ])
                    stats.rows += 1
            else:
                grouped[layer].append([
                    typ, handle, block, 0, None, None, None, None,
                    0, text_value, attrs, xdata,
                ])
                stats.rows += 1
                stats.skipped += 1

        if index == total or index % max(1, total // 100) == 0:
            progress(15 + int(index / total * 55), f"객체 분석 {index:,}/{total:,}")

    # ESSENPOLY는 원본 DXF 텍스트의 실제 레이어명으로 직접 출력한다.
    recovered_count = 0
    recovered_rows = 0
    per_layer = {}
    for item in recovered_items:
        layer = _norm(item.get("layer", ""))
        if layer not in grouped:
            continue
        points = item.get("points", [])
        if len(points) < 2:
            continue

        handle = _norm(item.get("handle", ""))
        meta = dict(item.get("attributes", {}))
        if item.get("color_aci") is not None:
            meta["CAD_COLOR_ACI"] = item.get("color_aci")
        if item.get("true_color") is not None:
            meta["CAD_TRUE_COLOR"] = item.get("true_color")
        meta["RECOVERED_TYPE"] = "ESSENPOLY/Embedded AcDbPolyline"
        attrs = _joined_attrs(meta)

        xmap = item.get("xdata", {}) or {}
        xdata = " | ".join(
            f"{appid}: " + " ; ".join(str(v) for v in values)
            for appid, values in xmap.items()
            if values
        )
        length = _length(points)

        entity_counts[layer] += 1
        stats.entities += 1
        recovered_count += 1
        per_layer[layer] = per_layer.get(layer, 0) + 1

        for seq, (x, y) in enumerate(points, 1):
            try:
                lon, lat = transformer.transform(x, y)
            except Exception:
                lon, lat = None, None
            grouped[layer].append([
                "ESSENPOLY", handle, "", seq, x, y, lon, lat,
                length, "", attrs, xdata,
            ])
            stats.rows += 1
            recovered_rows += 1

    if recovered_count:
        summary = ", ".join(f"{k}:{v}" for k, v in sorted(per_layer.items()))
        log(f"ESSENPOLY Excel 직접 복구 {recovered_count}개 / {recovered_rows}행 · {summary}")

    progress(78, "Excel 시트 생성")
    wb = Workbook()
    wb.remove(wb.active)
    used, index_rows = set(), []
    for layer_index, layer in enumerate(selected, 1):
        sheet_name = _safe_sheet_name(layer, used)
        ws = wb.create_sheet(sheet_name)
        _write_sheet(ws, LAYER_HEADERS, grouped[layer])
        index_rows.append([layer, sheet_name, entity_counts[layer], len(grouped[layer])])
        progress(78 + int(layer_index / max(1, len(selected)) * 14), f"시트 {layer_index}/{len(selected)}")

    idx = wb.create_sheet("LAYER_INDEX", 0)
    _write_sheet(idx, ["원본레이어","Excel시트","객체수","출력행수"], index_rows)
    info = wb.create_sheet("INFO")
    _write_sheet(info, ["항목","값"], [
        ["원본파일", str(input_path)],
        ["입력형식", "DXF only"],
        ["원본좌표계", f"EPSG:{source_epsg}"],
        ["출력좌표계", "EPSG:4326 (WGS84)"],
        ["선택레이어수", stats.layers],
        ["대상객체수", stats.entities],
        ["출력행수", stats.rows],
        ["ESSENPOLY복구객체수", recovered_count],
        ["ESSENPOLY복구행수", recovered_rows],
        ["외부전송", "없음 - 로컬 읽기 전용 Viewer/Excel 추출"],
    ])

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    progress(96, "Excel 저장")
    wb.save(out)
    progress(100, "완료")
    log(f"완료: {out}")
    return stats

def convert_file(input_path, output_path, source_epsg=5174, oda_exe=None, log=None, progress=None):
    layers = scan_layers(input_path, None, log, progress)
    return convert_selected_layers(
        input_path, output_path, [x.name for x in layers],
        source_epsg, None, log, progress
    )
