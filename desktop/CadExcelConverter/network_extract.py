from __future__ import annotations
import math
import re
from collections import deque
from dataclasses import dataclass
from types import SimpleNamespace
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pyproj import Transformer
from essenpoly_recovery import recover_essenpoly_polylines, recover_linker_polylines
from drawing_identity import export_object_id
from conduit_info import conduit_info, conduit_points


def xvalues(entity, appid, code):
    rows = dict(entity.xdata or []).get(appid, [])
    prefix = str(code) + ':'
    return [str(v).split(':', 1)[1].strip() for v in rows if str(v).startswith(prefix)]


def entity_position(entity):
    value = entity.dxf_data.get('insert', '')
    nums = re.findall(r'[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?', value)
    if len(nums) >= 2:
        return float(nums[0]), float(nums[1])
    for kind, data in entity.primitives:
        if kind in {'insert', 'point'}:
            return tuple(data[:2])
    if entity.bbox:
        x1,y1,x2,y2 = entity.bbox
        return (x1+x2)/2, (y1+y2)/2
    return None


def is_fiber(entity):
    name = entity.layer.upper()
    return any(k in name for k in ('F_CABLE', 'FOC', 'FIBER', '광케이블'))


def is_closure(entity):
    return entity.entity_type == 'INSERT' and ('CLOSURE' in entity.layer.upper() or '함체' in entity.layer)


def is_background(entity):
    name = str(entity.layer or '').upper()
    if is_fiber(entity) or any(word in name for word in ('CABLE', 'CONDUIT', 'PIPELINE', '관로', '케이블', '전주', '함체')):
        return False
    if name.startswith('CN_'):
        return name in {'CN_M_USER_BUILDING', 'CN_M_USER_REGION'}
    return True


def cable_length(points):
    return sum(math.dist(a,b) for a,b in zip(points,points[1:]))


@dataclass
class FiberEdge:
    handle: str
    u: str
    v: str
    points: list
    index: int
    attributes: dict


class FiberNetwork:
    def __init__(self, scene, input_path):
        self.scene = scene
        self.edges = {}
        self.positions = {}
        self.adj = {}
        self.by_handle = {e.handle:e for e in scene.entities}
        recovered = {r['handle']:r for r in recover_essenpoly_polylines(input_path)}
        self.excluded_stubs = 0
        self.geometry_nodes = []
        endpoint_join_tolerance = 5.0
        def coordinate_node(point):
            # 함체/NODELINK가 없어도 케이블 끝점끼리 가까우면 동일 연결점으로 본다.
            # 선 중간 교차는 보지 않고 시작/끝점에만 적용한다.
            for key, existing in self.geometry_nodes:
                if math.dist(point,existing) <= endpoint_join_tolerance:
                    return key
            key = 'XY' + str(len(self.geometry_nodes))
            self.geometry_nodes.append((key,point))
            return key
        for e in scene.entities:
            if not is_fiber(e):continue
            item = recovered.get(e.handle)
            attrs = (item or {}).get('attributes',e.attributes)
            xdata = (item or {}).get('xdata',dict(e.xdata or []))
            if 'STUB' in (e.layer + ' ' + str(attrs.get('ESSEN_301',''))).upper() or '지선' in e.layer or 'EXMAP_ONULINK' in xdata:
                self.excluded_stubs += 1
                continue
            points = (item or {}).get('points')
            if not points:
                points = next((p for kind,p in e.primitives if kind in {'polyline','line'}),[])
            if len(points)<2:continue
            links = [v.split(':',1)[1].strip().upper() for v in xdata.get('EXMAP_NODELINK',[]) if str(v).startswith('1005:')]
            links = [v for v in links if v and v != '0']
            if len(links)==2:
                u,v=links
            else:
                u,v=coordinate_node(points[0]),coordinate_node(points[-1])
            if u==v:continue
            edge=FiberEdge(e.handle,u,v,points,e.index,attrs)
            self.edges[e.handle]=edge
            for node,point in ((u,points[0]),(v,points[-1])):
                target=self.by_handle.get(node)
                self.positions[node]=(entity_position(target) if target else None) or tuple(point)
                self.adj.setdefault(node,[]).append((v if node==u else u,e.handle))

    def unique_route(self,start,end):
        if start==end:raise ValueError('시작점과 끝점이 같습니다. 다른 함체/끝점을 선택해주세요.')
        def find(skip=None):
            parents={start:None};q=deque([start])
            while q:
                n=q.popleft()
                if n==end:
                    route=[];nodes=[end]
                    while parents[n] is not None:
                        previous,handle=parents[n];route.append(handle);nodes.append(previous);n=previous
                    return list(reversed(route)),list(reversed(nodes))
                for to,handle in self.adj.get(n,[]):
                    if handle!=skip and to not in parents:
                        parents[to]=(n,handle);q.append(to)
            return None
        found=find()
        if not found:raise ValueError('두 점 사이에 연결된 광케이블 경로가 없습니다. 지선(STUB)은 제외됩니다.')
        route,nodes=found
        if any(find(handle) is not None for handle in route):
            raise ValueError('두 점 사이에 여러 광경로가 있습니다. 두 점만으로 주간선을 확정할 수 없어 추출하지 않습니다. 더 짧은 구간의 함체를 선택해주세요.')
        closures={e.index for node in nodes for e in [self.by_handle.get(node)] if e and is_closure(e)}
        # Exact coordinate association is only used for drawings without node handles.
        for node in nodes:
            if not node.startswith('XY'):continue
            closures.update(e.index for e in self.scene.entities if is_closure(e) and entity_position(e) and math.dist(entity_position(e),self.positions[node])<=5.0)
        return {'kind':'fiber','route':route,'nodes':nodes,'network':self,'indices':{self.edges[h].index for h in route}|closures}



def _online_attrs(props):
    raw = props.get('attributes') if isinstance(props, dict) else {}
    attrs = {}
    if isinstance(raw, dict):
        attrs.update({k:v for k,v in raw.items() if k != 'fields'})
        fields = raw.get('fields')
        if isinstance(fields, dict):
            attrs.update(fields)
    return attrs


def _online_xdata(props):
    raw = props.get('xdata') if isinstance(props, dict) else {}
    if isinstance(raw, dict):
        return list(raw.items())
    if isinstance(raw, list):
        return raw
    return []


def _online_geometry_lines(geometry):
    typ = str((geometry or {}).get('type') or '')
    coords = (geometry or {}).get('coordinates') or []
    if typ == 'LineString':
        return [coords]
    if typ == 'MultiLineString':
        return list(coords)
    return []


def _online_feature_identity(feature, region, part_index=0):
    props = feature.get('properties') or {}
    base = (
        feature.get('id')
        or props.get('source_handle')
        or props.get('entity_id')
        or props.get('regional_object_id')
        or props.get('object_id')
        or ''
    )
    if not base:
        base = f'geom-{abs(hash(str(feature.get("geometry") or {})))}'
    return f'{region}:{base}:{part_index}'


def _online_node_links(props):
    """Return the two CAD endpoint node references, when transported by the API."""
    raw = props.get('xdata') or _online_attrs(props).get('xdata') or {}
    values = raw.get('EXMAP_NODELINK', []) if isinstance(raw, dict) else []
    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, (list, tuple)) and len(item) >= 2 and item[0] == 'EXMAP_NODELINK':
                values = item[1]
                break
    if isinstance(values, str):
        values = [values]
    if not isinstance(values, (list, tuple)):
        return []
    links = [str(v).split(':', 1)[1].strip().upper() for v in values
             if str(v).startswith('1005:')]
    links = [v for v in links if v and v != '0']
    return links if len(links) == 2 else []


class OnlineFiberNetwork:
    """전체 권한 지역의 서버 FIBER GeoJSON을 하나의 경계 없는 광망으로 구성."""

    online = True

    def __init__(self, features, epsg=5174):
        self.scene = None
        self.edges = {}
        self.positions = {}
        self.adj = {}
        self.by_handle = {}
        self.excluded_stubs = 0
        self.source_regions = set()
        self._transform = Transformer.from_crs('EPSG:4326', f'EPSG:{int(epsg)}', always_xy=True)

        point_features = []
        for feature in features:
            props = feature.get('properties') or {}
            region = str(props.get('_cmb_region_code') or props.get('region_code') or '')
            if region:
                self.source_regions.add(region)
            geometry = feature.get('geometry') or {}
            if str(geometry.get('type') or '') != 'Point':
                continue
            coords = geometry.get('coordinates') or []
            if not (isinstance(coords, (list, tuple)) and len(coords) >= 2):
                continue
            try:
                x,y = self._transform.transform(float(coords[0]), float(coords[1]))
            except Exception:
                continue
            point_features.append((x,y,feature))

        # 함체가 없거나 행정구역 경계에서 심벌 공백이 있어도
        # 케이블의 실제 시작/끝점끼리 5m 이내면 동일 연결점으로 결합한다.
        # 선 중간 교차는 절대 연결하지 않는다.
        endpoint_join_tolerance = 5.0
        endpoint_nodes = []
        endpoint_links = {}
        def coord_key(point, link=None, region=''):
            # CAD node reference wins over geometry distance within the same region.
            # Separate regions can reuse CAD handles, so they are not globally keyed.
            link_key = (region, link) if link else None
            if link_key and link_key in endpoint_links:
                return endpoint_links[link_key]
            for key,existing in endpoint_nodes:
                if math.dist(point,existing) <= endpoint_join_tolerance:
                    if link_key:
                        endpoint_links[link_key] = key
                    return key
            key=f'XY:{len(endpoint_nodes)}'
            endpoint_nodes.append((key,tuple(point)))
            if link_key:
                endpoint_links[link_key] = key
            return key

        def nearest_point_feature(point, max_dist=1.5):
            best = None
            best_d = max_dist
            for px,py,feature in point_features:
                d = math.hypot(point[0]-px, point[1]-py)
                if d <= best_d:
                    best_d = d
                    best = (px,py,feature)
            return best

        edge_index = 0
        seen_geometry = set()
        for feature in features:
            props = feature.get('properties') or {}
            if str(props.get('group_id') or '').upper() != 'FIBER':
                continue
            region = str(props.get('_cmb_region_code') or props.get('region_code') or '')
            attrs = _online_attrs(props)
            layer = str(props.get('layer') or attrs.get('layer') or 'ONLINE_FIBER')
            entity_type = str(props.get('entity_type') or '')
            if entity_type.upper() in {'POINT','INSERT'}:
                continue

            for part_index, raw_line in enumerate(_online_geometry_lines(feature.get('geometry') or {})):
                points = []
                for coord in raw_line or []:
                    if not (isinstance(coord, (list, tuple)) and len(coord) >= 2):
                        continue
                    try:
                        points.append(self._transform.transform(float(coord[0]), float(coord[1])))
                    except Exception:
                        pass
                if len(points) < 2:
                    continue

                # 동일 형상이 지역 경계 양쪽 데이터에 중복 저장된 경우 한 번만 사용.
                signature = tuple((round(x,2),round(y,2)) for x,y in points)
                reverse_signature = tuple(reversed(signature))
                geom_key = min(signature, reverse_signature)
                if geom_key in seen_geometry:
                    continue
                seen_geometry.add(geom_key)

                links = _online_node_links(props)
                u = coord_key(points[0], links[0] if links else None, region)
                v = coord_key(points[-1], links[1] if links else None, region)
                if u == v:
                    continue

                handle = _online_feature_identity(feature, region, part_index)
                while handle in self.edges:
                    handle += '#'

                edge = FiberEdge(handle,u,v,points,edge_index,attrs)
                self.edges[handle] = edge
                edge_index += 1

                for node,point in ((u,points[0]),(v,points[-1])):
                    chosen = nearest_point_feature(point)
                    node_pos = (chosen[0],chosen[1]) if chosen else tuple(point)
                    self.positions.setdefault(node,node_pos)
                    self.adj.setdefault(node,[]).append((v if node==u else u,handle))

                    if node not in self.by_handle:
                        if chosen:
                            pprops = chosen[2].get('properties') or {}
                            pattrs = _online_attrs(pprops)
                            self.by_handle[node] = SimpleNamespace(
                                index=-1,
                                handle=node,
                                layer=str(pprops.get('layer') or 'ONLINE_FIBER_NODE'),
                                entity_type=str(pprops.get('entity_type') or 'INSERT'),
                                block_name=str(pprops.get('block_name') or ''),
                                text=str(pprops.get('text') or ''),
                                attributes=pattrs,
                                xdata=_online_xdata(pprops),
                                dxf_data={},
                                bbox=(node_pos[0],node_pos[1],node_pos[0],node_pos[1]),
                                primitives=[('point',node_pos)],
                            )
                        else:
                            self.by_handle[node] = SimpleNamespace(
                                index=-1, handle=node, layer='ONLINE_FIBER_NODE',
                                entity_type='POINT', block_name='', text='',
                                attributes={}, xdata=[], dxf_data={}, bbox=(node_pos[0],node_pos[1],node_pos[0],node_pos[1]),
                                primitives=[('point',node_pos)],
                            )

    def unique_route(self,start,end):
        if start == end:
            raise ValueError('시작점과 끝점이 같습니다. 다른 끝점을 선택해주세요.')
        parents = {start: None}
        q = deque([start])
        while q:
            node = q.popleft()
            if node == end:
                break
            for to,handle in self.adj.get(node,[]):
                if to not in parents:
                    parents[to] = (node,handle)
                    q.append(to)
        if end not in parents:
            raise ValueError('전체 온라인 광망에서도 두 점 사이의 연결 경로를 찾지 못했습니다.')
        route = []
        nodes = [end]
        node = end
        while parents[node] is not None:
            previous,handle = parents[node]
            route.append(handle)
            nodes.append(previous)
            node = previous
        route.reverse()
        nodes.reverse()
        return {
            'kind':'fiber',
            'route':route,
            'nodes':nodes,
            'network':self,
            'indices':{self.edges[h].index for h in route},
            'online':True,
        }


def _normalize_regions(raw):
    if isinstance(raw, dict):
        raw = raw.get('items') or raw.get('regions') or []
    result = []
    seen = set()
    for row in raw or []:
        if isinstance(row, str):
            region = row.strip()
        elif isinstance(row, dict):
            region = str(row.get('id') or row.get('region_code') or row.get('code') or '').strip()
        else:
            region = ''
        if region and region.casefold() not in seen:
            seen.add(region.casefold())
            result.append(region)
    return result


def fetch_all_online_fiber(client, regions=None, epsg=5174, progress=None):
    """선택한 권한 지역의 FIBER를 완전 수집 후 하나의 네트워크로 결합."""
    allowed_regions = _normalize_regions(client.regions() or [])
    if not allowed_regions:
        raise RuntimeError('온라인 광추출 가능한 권한 지역이 없습니다.')

    if regions is None:
        regions = list(allowed_regions)
    else:
        requested = [str(x or '').strip() for x in regions if str(x or '').strip()]
        allowed_map = {r.casefold(): r for r in allowed_regions}
        regions = [allowed_map[x.casefold()] for x in requested if x.casefold() in allowed_map]

    if not regions:
        raise RuntimeError('온라인 광추출 지역을 하나 이상 선택해주세요.')

    merged = {}
    region_stats = []
    total_regions = len(regions)
    limit = 5000
    max_depth = 14

    def feature_key(feature, region):
        props = feature.get('properties') or {}
        ident = (
            feature.get('id')
            or props.get('source_handle')
            or props.get('entity_id')
            or props.get('regional_object_id')
            or props.get('object_id')
            or str(feature.get('geometry') or {})
        )
        return f'{region}|{ident}'

    for region_no,region in enumerate(regions,1):
        if progress:
            progress(
                int(5 + (region_no-1) / max(total_regions,1) * 80),
                f'온라인 광망 수집 · {region_no}/{total_regions} · {region}',
            )

        meta = client.online_layers(region) or {}
        bounds = meta.get('bounds') if isinstance(meta, dict) else None
        if not isinstance(bounds, dict):
            region_stats.append((region,0,'bounds_missing'))
            continue
        try:
            west=float(bounds['west']); south=float(bounds['south'])
            east=float(bounds['east']); north=float(bounds['north'])
        except Exception:
            region_stats.append((region,0,'bounds_invalid'))
            continue

        def fetch_cell(w,s,e,n,depth=0):
            raw = client.online_objects(
                region,
                f'{w},{s},{e},{n}',
                limit=limit,
                groups='FIBER',
                timeout=120,
            ) or {}
            raw_features = list(raw.get('features') or [])
            features = [
                f for f in raw_features
                if str(((f.get('properties') or {}).get('group_id') or '')).upper() == 'FIBER'
            ]
            if len(raw_features) < limit:
                return features
            if depth >= max_depth:
                raise RuntimeError(
                    f'온라인 FIBER 전체수집 제한 도달 · {region} · depth={depth}'
                )
            mx=(w+e)/2.0; my=(s+n)/2.0
            cell_merged={}
            for cell in ((w,s,mx,my),(mx,s,e,my),(w,my,mx,n),(mx,my,e,n)):
                for feature in fetch_cell(*cell,depth+1):
                    cell_merged[feature_key(feature,region)] = feature
            return list(cell_merged.values())

        features = fetch_cell(west,south,east,north)
        for feature in features:
            props = dict(feature.get('properties') or {})
            props['_cmb_region_code'] = region
            feature = dict(feature)
            feature['properties'] = props
            merged[feature_key(feature,region)] = feature
        region_stats.append((region,len(features),'ok'))

    features = list(merged.values())
    if not features:
        raise RuntimeError('권한 지역에서 온라인 FIBER 데이터를 찾지 못했습니다.')

    if progress:
        progress(90,f'전체 광망 연결 구성 · FIBER 객체 {len(features):,}개')
    network = OnlineFiberNetwork(features,epsg=epsg)
    if not network.edges:
        raise RuntimeError('온라인 FIBER 선형 케이블을 구성하지 못했습니다.')
    if progress:
        progress(100,f'온라인 전체 광망 준비 · 케이블 {len(network.edges):,}개')
    return network, {
        'regions':len(regions),
        'features':len(features),
        'edges':len(network.edges),
        'nodes':len(network.positions),
        'region_stats':region_stats,
    }


def conduit_records(scene,input_path,diameters=None):
    diameters=diameters or {}
    entities=[e for e in scene.entities if 'CONDUIT' in e.layer.upper()]
    recovered={item['handle']:item for item in recover_linker_polylines(input_path)} if any(e.entity_type=='ASDKESSENLINKER' for e in entities) else {}
    records=[]
    for e in entities:
        item=recovered.get(e.handle)
        if item is None:
            item={'handle':e.handle,'layer':e.layer,'points':conduit_points(e),'attributes':e.attributes,'xdata':dict(e.xdata or [])}
        info=conduit_info(e,item)
        diameter=info['diameter'] if info['diameter'] is not None else diameters.get(e.handle)
        records.append({'entity':e,'item':item,'code':info['code'],'diameter':diameter,'count':info['count'],'raw':info['raw']})
    return records


def _node_name(network,node):
    e=network.by_handle.get(node)
    if e:
        values=xvalues(e,'EXMAP_FIBER',1000)
        if len(values)>1 and values[1]:return values[1]
        for key,value in e.attributes.items():
            if value and any(k in key.upper() for k in ('ID','NAME','번호')):return str(value)
    return node


def network_export_indices(scene, scope, visible_layers, shown, entity_items):
    if scope['kind']=='pipe':
        return {r['entity'].index for r in scope['records']
                if r['diameter']==100 and len(r['item']['points'])>=2}
    return {e.index for e in scene.entities if e.index in scope['indices']
            and (shown is None or e.index in shown) and e.layer in visible_layers
            and entity_items.get(e.index)}


def _online_export_ids(handle):
    """Convert an internal region:revision:CAD-handle:part key to stable, readable IDs."""
    value = str(handle or '').strip()
    region, sep, remainder = value.partition(':')
    if not sep:
        return value, value
    pieces = remainder.rsplit(':', 2)
    # Source identity is revision UUID : original CAD handle : geometry part.
    if len(pieces) == 3 and pieces[-1].isdigit():
        pure = pieces[-2]
        part = int(pieces[-1])
        if part:
            pure = f'{pure}:{part}'
    else:
        pure = remainder.rsplit(':', 1)[-1]
    return pure, f'{region}_{pure}'


def export_network(path,scene,scope,visible_indices,epsg=5174,source_path=None):
    visible_indices=set(visible_indices) & set(scope.get("indices", visible_indices))
    wb=Workbook();wb.remove(wb.active)
    rows=0
    geometry=[]
    def sheet(name,headers,data):
        nonlocal rows
        ws=wb.create_sheet(name)
        id_column=headers.index('객체ID')
        headers=list(headers)
        headers[id_column:id_column+1]=['순수 객체ID','지역객체ID']
        ws.append(headers)
        for row in data:
            if scope.get('online'):
                pure_id, regional_id = _online_export_ids(row[id_column])
                row[id_column] = pure_id
                row.insert(id_column+1, regional_id)
            else:
                row.insert(id_column+1,export_object_id(source_path,row[id_column]))
        for row in data:ws.append(row);rows+=1
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='176D85')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(42,max(14,max(len(str(c.value or '')) for c in col)+2))
    transform=Transformer.from_crs(f'EPSG:{epsg}','EPSG:4326',always_xy=True)
    if scope['kind']=='fiber':
        network=scope['network'];data=[]
        routes=list(scope.get('routes') or [])
        if not routes:
            routes=[{'route':scope.get('route') or [],'nodes':scope.get('nodes') or []}]
        exported_handles=set()
        for route_no,route_scope in enumerate(routes,1):
            route=list(route_scope.get('route') or [])
            nodes=list(route_scope.get('nodes') or [])
            for order,handle in enumerate(route,1):
                edge=network.edges[handle]
                if edge.index not in visible_indices or handle in exported_handles:
                    continue
                exported_handles.add(handle)
                a,b=nodes[order-1:order+1]
                summary=str(edge.attributes.get('ESSEN_300','')).split('/')
                cable_id=edge.attributes.get('ESSEN_302') or summary[0] or handle
                cores=summary[1] if len(summary)>1 else ''
                declared=None
                if len(summary)>2:
                    try:declared=float(summary[2])
                    except ValueError:pass
                points = edge.points if edge.u == a else list(reversed(edge.points))
                start_point, end_point = points[0], points[-1]
                if scope.get('online'):
                    # Use one canonical coordinate per graph node so consecutive
                    # rows are exactly continuous even when source endpoints differ.
                    start_point = network.positions.get(a, start_point)
                    end_point = network.positions.get(b, end_point)
                start_closure = network.by_handle.get(a)
                end_closure = network.by_handle.get(b)
                if order > 1 and start_closure and is_closure(start_closure):
                    start_point = entity_position(start_closure) or start_point
                if order < len(route) and end_closure and is_closure(end_closure):
                    end_point = entity_position(end_closure) or end_point
                start_lon, start_lat = transform.transform(*start_point)
                end_lon, end_lat = transform.transform(*end_point)
                data.append([route_no,order,cable_id,cores,declared,round(cable_length(edge.points),3),_node_name(network,a),_node_name(network,b),handle,start_lon,start_lat,end_lon,end_lat])
        if not data:raise ValueError('화면에 선택된 광케이블이 없습니다.')
        sheet('광케이블',['경로','순서','케이블번호','심수','도면기재길이_m','경로길이_m','시작함체','끝함체','객체ID','시작경도','시작위도','끝경도','끝위도'],data)
        ws=wb['광케이블']
        for row in ws.iter_rows(min_row=2,min_col=11,max_col=14):
            for cell in row:cell.number_format='0.0000000'
    else:
        from pipe_endpoints import PipeEndpointResolver
        resolver=PipeEndpointResolver(scene)
        data=[]
        for r in scope['records']:
            if r['diameter']!=100 or r['entity'].index not in visible_indices or len(r['item']['points'])<2:continue
            pts=r['item']['points']
            start=resolver.resolve(pts[0]);end=resolver.resolve(pts[-1])
            start_lon,start_lat=transform.transform(*start['point'])
            end_lon,end_lat=transform.transform(*end['point'])
            data.append([r['entity'].handle,100,r.get('count'),r.get('raw',''),round(cable_length(pts),3),start['type'],start['name'],start_lon,start_lat,end['type'],end['name'],end_lon,end_lat])
        if not data:raise ValueError('추출 가능한 100mm 주관로가 없습니다. 지름 미확인/50mm 관로는 추출하지 않습니다.')
        sheet('100mm_주관로',['객체ID','지름_mm','본수','원본관경표기','길이_m','시작시설','시작시설명_전주코드','시작경도','시작위도','끝시설','끝시설명_전주코드','끝경도','끝위도'],data)
        for row in wb['100mm_주관로'].iter_rows(min_row=2):
            for col in (9,10,13,14):row[col-1].number_format='0.0000000'
    wb.properties.description = f"CAD 좌표계 EPSG:{epsg}; 경위도 EPSG:4326"
    wb.save(path)
    return rows
