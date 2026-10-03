import tempfile
from pathlib import Path

import ezdxf
from openpyxl import load_workbook

from converter import convert_file

def main():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        dxf = root / "sample.dxf"
        xlsx = root / "sample.xlsx"

        doc = ezdxf.new("R2018")
        msp = doc.modelspace()

        pole_block = doc.blocks.new("Pole-Joint")
        pole_block.add_circle((0, 0), 1)
        msp.add_blockref(
            "Pole-Joint",
            (156215.0863507305, 174095.3802128593),
            dxfattribs={"layer": "CN_L_Pole_Pole-Joint"},
        )

        onu_block = doc.blocks.new("ONU")
        onu_block.add_circle((0, 0), 1)
        msp.add_blockref(
            "ONU",
            (156250.0, 174120.0),
            dxfattribs={"layer": "CN_C_ONU"},
        )

        msp.add_lwpolyline(
            [(156215.0, 174095.0), (156225.0, 174105.0), (156235.0, 174115.0)],
            dxfattribs={"layer": "CN_F_Cable_FOC"},
        )
        msp.add_lwpolyline(
            [(156240.0, 174100.0), (156250.0, 174110.0)],
            dxfattribs={"layer": "CN_C_Cable_500F"},
        )
        doc.saveas(dxf)

        stats = convert_file(str(dxf), str(xlsx), 5174)
        assert stats.facilities == 1, stats
        assert stats.equipment == 1, stats
        assert stats.fiber_points == 3, stats
        assert stats.coax_points == 2, stats

        wb = load_workbook(xlsx, data_only=True)
        assert ["CELL", "FACILITY", "EQUIPMENT", "FIBER", "COAX", "INFO"] == wb.sheetnames
        ws = wb["FACILITY"]
        lon = ws.cell(2, 7).value
        lat = ws.cell(2, 8).value
        assert abs(lon - 126.5208008) < 0.0001, lon
        assert abs(lat - 35.0649176) < 0.0001, lat

        print("SELFTEST OK")

if __name__ == "__main__":
    main()
