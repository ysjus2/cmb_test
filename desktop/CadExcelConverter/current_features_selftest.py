from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from openpyxl import load_workbook
from pyproj import Transformer
from main import App
from converter import LayerInfo
from viewer import Scene, VisualEntity, _optical_arrow_primitive
from conduit_info import conduit_info, conduit_points
from network_extract import conduit_records, export_network, network_export_indices
from catv_poles import catv_pole_rows, export_catv_poles, is_catv_pole

def pipe(index, handle, raw, x):
    return VisualEntity(index, 'ACAD_PROXY_ENTITY', 'CN_L_Pole_Line_Conduit', handle,
        primitives=[('text', (x, 0, handle)), ('polyline', [(x, 0), (x+5, 0)]),
                    ('polyline', [(x+5, 0), (x+10, 0)])], bbox=(x, 0, x+10, 0),
        xdata=[('EXMAP_POLELINK', ['1000: U', '1000: '+raw])])

entities = [pipe(0,'P1','100*1',0),pipe(1,'P2','100*3',10),pipe(2,'P3','50*1',20),
    VisualEntity(3,'LINE','TL_SPRD_RW','ROAD',primitives=[('line',[(0,1),(30,1)])],bbox=(0,1,30,1)),
    VisualEntity(4,'INSERT','CN_L_Pole_Pole-CATV','POLE',block_name='Pole-CATV',attributes={'ID':'7887X123'},dxf_data={'insert':'10.01,0,0'},primitives=[('circle',(10.01,0,1)),('text',(10.01,0,'pole name'))],bbox=(9,-1,12,1)),
    VisualEntity(5,'INSERT','CN_L_Pole_Manhole-CATV','MH',block_name='Manhole-CATV',attributes={'ID':'MH1'},dxf_data={'insert':'0.01,0,0'},primitives=[('circle',(.01,0,1))],bbox=(-1,-1,1,1))]
scene = Scene(entities,(-1,-1,31,2),{})
assert conduit_info(entities[1])['count']==3
assert conduit_points(entities[0])==[(0,0),(5,0),(10,0)]
assert not is_catv_pole(entities[5])
assert is_catv_pole(VisualEntity(0,'INSERT','CN_L_Pole_Pole_CATV','A'))
arrows=[('line',[(0,0),(1,0)]),('line',[(1,0),(2,0)]),('polyline',[(2,0),(3,1),(2,0)])]
assert _optical_arrow_primitive(arrows)==1
app=App();app.input_path='CMB_GN_DY.dxf'
app._populate_layers([LayerInfo(name,0,'') for name in sorted({e.layer for e in entities})],scene)
app.viewer.load_scene(scene);app.update()
scope={'kind':'pipe','records':conduit_records(scene,'unused.dxf'),'indices':{0,1}}
assert network_export_indices(scene,scope,set(),set(),{})=={0,1}
app._show_scope(scope)
assert {i for i,ids in app.viewer.entity_items.items() if ids}=={0,1}
assert not app.viewer.canvas.find_withtag('annotation')
assert len(list(app.viewer._display_entities()))==2
with TemporaryDirectory() as temporary:
    out=Path(temporary)/'pipe.xlsx'
    assert export_network(out,scene,scope,{0,1},source_path=app.input_path)==2
    ws=load_workbook(out)['100mm_주관로'];headers={cell.value:cell.column for cell in ws[1]}
    get=lambda name:ws.cell(2,headers[name]).value
    assert get('순수 객체ID')=='P1' and get('지역객체ID')=='CMB_GN_DY_P1'
    assert get('시작시설')=='맨홀' and get('끝시설')=='전주'
    transform=Transformer.from_crs('EPSG:5174','EPSG:4326',always_xy=True)
    assert abs(get('끝경도')-transform.transform(10.01,0)[0])<1e-9
    with patch('network_extract_ui.filedialog.asksaveasfilename',return_value=''):
        app._run_catv_pole_excel()
    assert {i for i,ids in app.viewer.entity_items.items() if ids}=={4}
    assert not app.viewer.canvas.find_withtag('annotation')
    rows,missing=catv_pole_rows(scene,app.input_path)
    assert len(rows)==1 and not missing
    assert export_catv_poles(Path(temporary)/'poles.xlsx',rows)==1
app._reset_network_filter()
assert app.viewer.entity_filter is None and not app.viewer.hide_extraction_text
app.destroy()
print('CURRENT FEATURES OK: target-only display, whole pipes, facility positions, CATV exclusion, dual IDs, single arrow')
