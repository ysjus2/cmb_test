import tempfile
from pathlib import Path

import ezdxf
from openpyxl import load_workbook

from converter import scan_layers, convert_selected_layers, load_dxf_document
from viewer import build_scene


def assert_has_primitive(scene, kind):
    for entity in scene.entities:
        for primitive in entity.primitives:
            if primitive[0] == kind:
                return
    raise AssertionError(f"missing primitive: {kind}")


def main():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        dxf = root / "functional.dxf"
        xlsx = root / "functional.xlsx"

        doc = ezdxf.new("R2018")
        msp = doc.modelspace()

        # Representative CMB geometry / annotation / equipment objects.
        msp.add_line((156200, 174090), (156210, 174100), dxfattribs={"layer": "LINE_TEST"})
        msp.add_lwpolyline(
            [(156210, 174100), (156220, 174110), (156230, 174120)],
            dxfattribs={"layer": "CN_F_Cable_FOC"},
        )
        msp.add_circle((156240, 174110), 5, dxfattribs={"layer": "CIRCLE_TEST"})
        msp.add_arc((156250, 174110), 7, 10, 170, dxfattribs={"layer": "ARC_TEST"})
        msp.add_point((156260, 174115), dxfattribs={"layer": "POINT_TEST"})
        msp.add_text("CMB-TEXT", dxfattribs={"layer": "ANNO"}).set_placement((156270, 174120))
        msp.add_mtext("CMB-MTEXT", dxfattribs={"layer": "ANNO"}).set_location((156275, 174125))

        pole = doc.blocks.new("Pole-Joint")
        pole.add_circle((0, 0), 1)
        pole.add_line((-2, 0), (2, 0))
        msp.add_blockref(
            "Pole-Joint", (156215.0863507305, 174095.3802128593),
            dxfattribs={"layer": "CN_L_Pole_Pole-Joint"},
        )

        doc.saveas(dxf)

        # DXF parser + layer scan.
        loaded = load_dxf_document(str(dxf))
        assert len(list(loaded.modelspace())) == 8
        layers = scan_layers(str(dxf))
        layer_names = {x.name for x in layers}
        for required in {
            "LINE_TEST", "CN_F_Cable_FOC", "CIRCLE_TEST", "ARC_TEST",
            "POINT_TEST", "ANNO", "CN_L_Pole_Pole-Joint"
        }:
            assert required in layer_names, required

        # Viewer scene build and representative primitives.
        scene = build_scene(str(dxf))
        assert len(scene.entities) == 8, len(scene.entities)
        assert scene.bbox[0] < scene.bbox[2]
        assert scene.bbox[1] < scene.bbox[3]
        for kind in ("line", "polyline", "circle", "arc", "point", "text", "insert"):
            assert_has_primitive(scene, kind)

        # Excel export + coordinate conversion.
        selected = ["CN_L_Pole_Pole-Joint", "CN_F_Cable_FOC", "ANNO"]
        stats = convert_selected_layers(str(dxf), str(xlsx), selected, 5174)
        assert stats.layers == 3
        assert stats.entities == 4  # pole + cable + TEXT + MTEXT
        assert stats.rows == 6      # pole 1 + cable vertices 3 + TEXT 1 + MTEXT 1

        wb = load_workbook(xlsx, data_only=True)
        assert wb.sheetnames[0] == "LAYER_INDEX"
        assert "INFO" in wb.sheetnames
        idx = wb["LAYER_INDEX"]
        mapping = {idx.cell(r, 1).value: idx.cell(r, 2).value for r in range(2, idx.max_row + 1)}
        assert set(mapping) == set(selected)

        pole_ws = wb[mapping["CN_L_Pole_Pole-Joint"]]
        lon = pole_ws.cell(2, 7).value
        lat = pole_ws.cell(2, 8).value
        assert abs(lon - 126.5208008) < 0.0001, lon
        assert abs(lat - 35.0649176) < 0.0001, lat

        fiber_ws = wb[mapping["CN_F_Cable_FOC"]]
        assert fiber_ws.max_row == 4
        assert fiber_ws.cell(2, 1).value == "LWPOLYLINE"

        anno_ws = wb[mapping["ANNO"]]
        texts = {anno_ws.cell(r, 10).value for r in range(2, anno_ws.max_row + 1)}
        assert "CMB-TEXT" in texts
        assert "CMB-MTEXT" in texts

        # Non-DXF input must remain rejected.
        bad = root / "bad.txt"
        bad.write_text("not dxf", encoding="utf-8")
        try:
            load_dxf_document(str(bad))
        except ValueError:
            pass
        else:
            raise AssertionError("non-DXF file was not rejected")

        print("FUNCTIONAL REGRESSION OK")


if __name__ == "__main__":
    main()
