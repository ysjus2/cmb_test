from pathlib import Path
import tempfile

from openpyxl import load_workbook

from essenpoly_recovery import recover_essenpoly_polylines
from converter import convert_selected_layers


def main():
    sample = """  0\nSECTION\n  2\nENTITIES\n  0\nESSENPOLY\n  5\nABC1\n100\nAcDbEntity\n  8\nCN_F_Cable_FOC\n300\nCABLE-A\n101\nEmbedded Object\n100\nAcDbEntity\n  8\nCN_F_Cable_FOC\n 62\n5\n100\nAcDbPolyline\n 90\n        3\n 70\n     0\n 10\n200000.0\n 20\n500000.0\n 10\n200010.0\n 20\n500005.0\n 10\n200020.0\n 20\n500010.0\n1001\nEXMAP_NODELINK\n1000\nNODE-A\n1000\nNODE-B\n  0\nENDSEC\n  0\nEOF\n"""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.dxf"
        xlsx = Path(td) / "sample.xlsx"
        path.write_text(sample, encoding="utf-8")

        rows = recover_essenpoly_polylines(path)
        assert len(rows) == 1, rows
        row = rows[0]
        assert row["handle"] == "ABC1", row
        assert row["layer"] == "CN_F_Cable_FOC", row
        assert row["points"] == [
            (200000.0, 500000.0),
            (200010.0, 500005.0),
            (200020.0, 500010.0),
        ], row
        assert row["attributes"]["ESSEN_300"] == "CABLE-A", row
        assert row["color_aci"] == 5, row

        stats = convert_selected_layers(
            path,
            xlsx,
            ["CN_F_Cable_FOC"],
            source_epsg=5174,
        )
        assert stats.entities == 1, stats
        assert stats.rows == 3, stats

        wb = load_workbook(xlsx, data_only=True)
        ws = wb["CN_F_Cable_FOC"]
        assert ws.max_row == 4, ws.max_row
        assert ws["A2"].value == "ESSENPOLY", ws["A2"].value
        assert ws["B2"].value == "ABC1", ws["B2"].value
        assert ws["E2"].value == 1, ws["E2"].value
        assert ws["F2"].value == 200000.0, ws["F2"].value
        assert ws["G2"].value == 500000.0, ws["G2"].value
        assert "CABLE-A" in str(ws["L2"].value), ws["L2"].value

        idx = wb["LAYER_INDEX"]
        assert idx["C2"].value == 1, idx["C2"].value
        assert idx["D2"].value == 3, idx["D2"].value

    print("ESSENPOLY recovery + Excel export self-test OK")


if __name__ == "__main__":
    main()
