from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from pyproj import Transformer
from viewer import VisualEntity,Scene
from network_extract import export_network
from pipe_endpoints import PipeEndpointResolver,facility_name


def main():
    pipe=VisualEntity(index=0,entity_type='ASDKESSENLINKER',layer='CN_L_Pole_Line_Conduit',handle='AB',primitives=[('polyline',[(0,0),(4,8),(10,0)])])
    mh=VisualEntity(index=1,entity_type='INSERT',layer='CN_L_Pole_Manhole-PWR',handle='MH',attributes={'ID':'맨홀1'},dxf_data={'insert':'0, 0, 0'})
    pole=VisualEntity(index=2,entity_type='INSERT',layer='CN_L_Pole_Pole-Joint',handle='POLE',attributes={'ID':'7887X123'},dxf_data={'insert':'10, 0, 0'})
    scene=Scene([pipe,mh,pole],(0,0,10,8),{})
    records=[{'entity':pipe,'diameter':100,'item':{'points':[(0,0),(4,8),(10,0)]}}]
    scope={'kind':'pipe','records':records,'indices':{0}}
    resolver=PipeEndpointResolver(scene)
    assert resolver.resolve((0,0))['type']=='맨홀'
    assert resolver.resolve((10,0))['name']=='7887X123'
    assert resolver.resolve((10,0.2))['type']=='미확인'
    unknown=resolver.resolve((100,100));assert unknown['type']=='미확인' and unknown['point']==(100,100)
    pole.attributes={};pole.xdata=[('EXMAP_POLE',['1000: 7887X123'])]
    assert facility_name(pole)=='7887X123'
    with TemporaryDirectory() as td:
        out=Path(td)/'pipes.xlsx'
        rows=export_network(out,scene,scope,{0})
        wb=load_workbook(out,data_only=True)
        assert wb.sheetnames==['100mm_주관로'] and rows==1
        ws=wb['100mm_주관로']
        assert ws.max_row==2 and ws['D2'].value=='맨홀' and ws['E2'].value=='맨홀1'
        assert ws['H2'].value=='전주' and ws['I2'].value=='7887X123'
        transform=Transformer.from_crs('EPSG:5174','EPSG:4326',always_xy=True)
        assert abs(ws['F2'].value-transform.transform(0,0)[0])<1e-9
        assert abs(ws['K2'].value-transform.transform(10,0)[1])<1e-9
    scene.entities.append(VisualEntity(index=3,entity_type='INSERT',layer='CN_L_Pole_Manhole-Drop',handle='MH2',dxf_data={'insert':'0, 0, 0'}))
    assert PipeEndpointResolver(scene).resolve((0,0))['type']=='복수시설'
    print('Pipe endpoint test OK: manhole GPS, endpoint pole code, hidden-node lookup, single sheet, missing/overlapping nodes flagged')

if __name__=='__main__':main()
