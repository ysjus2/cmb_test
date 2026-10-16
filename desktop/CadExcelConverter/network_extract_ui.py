from __future__ import annotations
import json
import math
import threading
import tkinter as tk
from pathlib import Path
from tkinter import ttk, messagebox, filedialog
from converter import ConversionStats
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pyproj import Transformer
from layer_defaults import is_default_hidden
from network_extract import FiberNetwork, fetch_all_online_fiber, conduit_records, export_network, export_submission_network, is_fiber, is_closure, network_export_indices
from catv_poles import catv_pole_rows, export_catv_poles


class NetworkExtractionMixin:
    def _open_fast_search(self):
        """Search only the current scene, or the current extracted network."""
        from network_extract import entity_position, _online_export_ids
        from pyproj import Transformer
        import re
        if getattr(self,'_search_window',None) is not None:
            try:
                if self._search_window.winfo_exists():
                    self._search_window.lift()
                    return
            except Exception:
                pass
        scope = getattr(self,'export_scope',None)
        network = (scope or {}).get('network') if scope and scope.get('kind')=='fiber' else None
        entries = []
        if network is not None:
            allowed = set((scope or {}).get('indices') or [])
            for handle,edge in network.edges.items():
                if edge.index not in allowed:
                    continue
                pts = list(edge.points or [])
                if not pts:
                    continue
                attrs = edge.attributes or {}
                cable = str(attrs.get('ESSEN_302') or
                            str(attrs.get('ESSEN_300') or '').split('/')[0]).strip()
                pure, regional = _online_export_ids(handle) if getattr(network,'online',False) else (handle,handle)
                entries.append((str(pure),str(regional),cable,tuple(pts[len(pts)//2]),edge.index,pts))
        elif getattr(self.viewer,'scene',None) is not None:
            for entity in self.viewer.scene.entities:
                attrs = entity.attributes or {}
                cable = str(attrs.get('ESSEN_302') or
                            str(attrs.get('ESSEN_300') or '').split('/')[0]).strip()
                points = next((list(v) for k,v in entity.primitives
                               if k in ('line','polyline') and len(v)>=2),[])
                position = entity_position(entity)
                if position is None and points:
                    position = points[len(points)//2]
                if position is not None:
                    entries.append((str(entity.handle),str(entity.handle),cable,
                                    tuple(position[:2]),entity.index,points))
        if not entries:
            messagebox.showinfo('빠른 검색','현재 검색할 도면 또는 추출 데이터가 없습니다.')
            return
        win=tk.Toplevel(self)
        self._search_window=win
        win.title('현재 도면 / 광추출 빠른 검색')
        win.geometry('550x390')
        win.transient(self)
        pane=ttk.Frame(win,padding=10)
        pane.pack(fill='both',expand=True)
        mode=tk.StringVar(value='객체ID')
        term=tk.StringVar()
        top=ttk.Frame(pane)
        top.pack(fill='x')
        ttk.Combobox(top,textvariable=mode,values=('객체ID','좌표','케이블번호'),state='readonly',width=11).pack(side='left',padx=(0,4))
        box=ttk.Entry(top,textvariable=term)
        box.pack(side='left',fill='x',expand=True)
        table=ttk.Treeview(pane,columns=('id','cable'),show='headings',height=12)
        table.heading('id',text='객체ID')
        table.heading('cable',text='케이블번호')
        table.column('id',width=330)
        table.column('cable',width=125)
        table.pack(fill='both',expand=True,pady=(10,3))
        info=ttk.Label(pane,text=f'검색 대상 {len(entries):,}개 · 현재 데이터만 조회')
        info.pack(anchor='w')
        matches=[]
        transformer=Transformer.from_crs('EPSG:4326',f'EPSG:{int(self.epsg_var.get())}',always_xy=True)

        def search(*_):
            nonlocal matches
            query=term.get().strip()
            table.delete(*table.get_children())
            if not query:
                matches=[]
                return
            option=mode.get()
            if option=='좌표':
                try:
                    items=[float(v) for v in re.split(r'[,\s]+',query) if v]
                    if len(items)!=2 or not (-180<=items[0]<=180 and -90<=items[1]<=90):
                        raise ValueError()
                    x,y=transformer.transform(*items)
                    ranked=sorted(entries,key=lambda e:(e[3][0]-x)**2+(e[3][1]-y)**2)
                    matches=ranked[:50]
                except (ValueError,TypeError):
                    info.configure(text='경도, 위도 순서로 입력해주세요.')
                    return
            elif option=='케이블번호':
                matches=[e for e in entries if query.upper() in e[2].upper()][:100]
            else:
                matches=[e for e in entries if query.upper() in e[0].upper() or query.upper() in e[1].upper()][:100]
            for i,e in enumerate(matches):
                table.insert('', 'end', iid=str(i),values=(e[1],e[2]))
            info.configure(text=f'검색 결과 {len(matches):,}개 · 더블클릭 또는 Enter로 이동')

        def move(_=None):
            choice=table.selection()
            if not choice:
                return
            e=matches[int(choice[0])]
            x,y=e[3]
            viewer=self.viewer
            viewer.update_idletasks()
            w=max(100,viewer.canvas.winfo_width())
            h=max(100,viewer.canvas.winfo_height())
            viewer.ox=w/2-x*viewer.scale
            viewer.oy=h/2+y*viewer.scale
            viewer.selected.clear()
            viewer.selected.add(e[4])
            viewer.redraw()
            info.configure(text=f'선택: {e[1]} · 도면 중앙으로 이동')
        ttk.Button(top,text='검색',command=search).pack(side='left',padx=(4,0))
        box.bind('<Return>',search)
        table.bind('<Double-1>',move)
        table.bind('<Return>',move)
        box.focus_set()

    def _build_extract_toolbar(self,parent,fiber_parent=None):
        self.export_scope=None
        self.fiber_network=None
        self.route_start=None
        self.fiber_routes=[]
        self.fiber_history=[]
        self.fiber_pick_mode=None
        self.pipe_diameters={}
        self.extract_bar=parent

        can_extract=bool(getattr(self,'can_extract',False))
        if can_extract:
            self.extract_menu_btn=ttk.Button(
                self.extract_bar,
                text='위·경도 좌표 추출  ▼',
                command=self._show_extract_menu,
            )
            self.extract_menu_btn.pack(side='left',padx=(18,6))
            self.extract_menu=tk.Menu(self.extract_bar,tearoff=False)
            self.extract_menu.add_command(label='광케이블 추출',command=self._begin_fiber_route)
            self.extract_menu.add_command(label='100mm 주관로 추출',command=self._begin_main_conduit)
            self.extract_menu.add_command(label='자가주 좌표 추출',command=self._run_catv_pole_excel)
            self.extract_menu.add_command(label='사용자 지정 영역 추출',command=self._begin_custom_area_extract)
            self.extract_menu.add_separator()
            self.extract_menu.add_command(label='Excel 추출',command=self._run_excel)
            self.extract_menu.add_command(label='광선로 제출용 엑셀 (8열)',command=lambda:self._run_network_excel(submission=True))
            self.extract_menu.add_separator()
            self.extract_menu.add_command(label='온라인 광추출',command=self._begin_online_fiber_route)

        if bool(getattr(self,'is_admin',False)):
            ttk.Button(
                self.extract_bar,
                text='도면 업로드',
                command=self._open_admin_upload,
            ).pack(side='left',padx=(24,6))

        self.fiber_controls=ttk.Frame(fiber_parent or parent)
        self.fiber_start_btn=ttk.Button(
            self.fiber_controls,text='시작점 선택',
            command=lambda:self._set_fiber_pick_mode('start'),
        )
        self.fiber_end_btn=ttk.Button(
            self.fiber_controls,text='끝점 선택',
            command=lambda:self._set_fiber_pick_mode('end'),
        )
        self.fiber_undo_btn=ttk.Button(
            self.fiber_controls,text='한 단계 취소',
            command=self._undo_fiber_selection,
        )
        self.fiber_clear_btn=ttk.Button(
            self.fiber_controls,text='전체 초기화',
            command=self._clear_fiber_selections,
        )
        self.fiber_excel_btn=ttk.Button(
            self.fiber_controls,text='Excel 추출',
            command=self._run_excel,
        )
        for button in (
            self.fiber_start_btn,self.fiber_end_btn,self.fiber_undo_btn,
            self.fiber_clear_btn,self.fiber_excel_btn,
        ):
            button.pack(side='left',padx=(0,5))

        # 상태 문구는 내부 동작용으로만 유지하고 상단에는 표시하지 않는다.
        hint=(
            '위·경도 좌표 추출 메뉴에서 작업을 선택하세요.'
            if can_extract else
            '보기 전용 계정 · 추출/Excel 기능 없음'
        )
        self.extract_hint=tk.StringVar(value=hint)

    def _begin_online_fiber_route(self):
        if not getattr(self,'can_extract',False):
            return
        if self.busy:
            return

        client=self._get_session_client()
        try:
            raw=client.regions() or []
        except Exception as exc:
            messagebox.showerror('온라인 광추출',f'지역 목록 조회 실패\n{exc}')
            return

        rows=raw.get('items') or raw.get('regions') or [] if isinstance(raw,dict) else raw
        regions=[]
        for row in rows or []:
            if isinstance(row,str):
                code=row.strip()
                name=code
            elif isinstance(row,dict):
                code=str(row.get('id') or row.get('region_code') or row.get('code') or '').strip()
                name=str(row.get('name') or code).strip()
            else:
                continue
            if code:
                regions.append((code,name or code))

        if not regions:
            messagebox.showinfo('온라인 광추출','선택 가능한 권한 지역이 없습니다.')
            return

        win=tk.Toplevel(self)
        win.title('온라인 광추출 · 지역 선택')
        win.geometry('520x560')
        win.minsize(460,420)
        win.transient(self)
        win.grab_set()

        root=ttk.Frame(win,padding=12)
        root.pack(fill='both',expand=True)

        ttk.Label(
            root,
            text='온라인 광추출 지역 선택',
            font=('Malgun Gothic',14,'bold'),
        ).pack(anchor='w')
        ttk.Label(
            root,
            text='광경로가 통과할 수 있는 지역을 모두 선택하세요. 복수 선택 가능합니다.',
        ).pack(anchor='w',pady=(3,10))

        box=ttk.LabelFrame(root,text='권한 지역',padding=8)
        box.pack(fill='both',expand=True)

        canvas=tk.Canvas(box,highlightthickness=0)
        scroll=ttk.Scrollbar(box,orient='vertical',command=canvas.yview)
        body=ttk.Frame(canvas)
        body_win=canvas.create_window((0,0),window=body,anchor='nw')
        body.bind('<Configure>',lambda e:canvas.configure(scrollregion=canvas.bbox('all')))
        canvas.bind('<Configure>',lambda e:canvas.itemconfigure(body_win,width=e.width))
        canvas.configure(yscrollcommand=scroll.set)
        canvas.pack(side='left',fill='both',expand=True)
        scroll.pack(side='right',fill='y')

        vars_by_region={}
        for code,name in regions:
            var=tk.BooleanVar(value=True)
            vars_by_region[code]=var
            label=f'{name}  ({code})' if name!=code else code
            ttk.Checkbutton(body,text=label,variable=var).pack(anchor='w',pady=2)

        def set_all(value):
            for var in vars_by_region.values():
                var.set(bool(value))

        def load_selected():
            selected=[code for code,var in vars_by_region.items() if var.get()]
            if not selected:
                messagebox.showinfo(
                    '온라인 광추출',
                    '지역을 하나 이상 선택해주세요.',
                    parent=win,
                )
                return

            win.destroy()
            self.busy=True
            self._show_progress()
            self._set_progress(3,f'온라인 광망 준비 · 지역 {len(selected)}개')
            self.viewer.status_var.set(
                f'온라인 광추출 · 선택 지역 {len(selected)}개 불러오는 중'
            )

            def worker():
                try:
                    network,summary=fetch_all_online_fiber(
                        client,
                        regions=selected,
                        epsg=int(self.epsg_var.get()),
                        progress=lambda p,t:self.q.put(('progress',(p,t))),
                    )
                    self.q.put(('online_fiber_loaded',(network,summary,selected)))
                except Exception as exc:
                    self.q.put(('error',str(exc)))

            threading.Thread(target=worker,daemon=True).start()

        footer=ttk.Frame(root)
        footer.pack(fill='x',pady=(10,0))
        ttk.Button(footer,text='전체 선택',command=lambda:set_all(True)).pack(side='left')
        ttk.Button(footer,text='전체 해제',command=lambda:set_all(False)).pack(side='left',padx=(6,0))
        ttk.Button(footer,text='불러오기',command=load_selected).pack(side='right')
        ttk.Button(footer,text='취소',command=win.destroy).pack(side='right',padx=(0,6))

    def _show_extract_menu(self):
        menu=getattr(self,'extract_menu',None)
        button=getattr(self,'extract_menu_btn',None)
        if menu is None or button is None:
            return
        try:
            x=button.winfo_rootx()
            y=button.winfo_rooty()+button.winfo_height()
            menu.tk_popup(x,y)
        finally:
            try:
                menu.grab_release()
            except Exception:
                pass

    def _show_fiber_controls(self,show=True):
        host=getattr(self,'viewer',None)
        host=getattr(host,'fiber_controls_host',None)
        if show:
            if host is not None and not host.winfo_manager():
                host.pack(fill='x',pady=(0,4),before=self.viewer.body)
            if not self.fiber_controls.winfo_manager():
                self.fiber_controls.pack(fill='x')
        else:
            if self.fiber_controls.winfo_manager():
                self.fiber_controls.pack_forget()
            if host is not None and host.winfo_manager():
                host.pack_forget()

    def _update_fiber_pick_buttons(self):
        fixed = self.route_start is not None
        try:
            self.fiber_start_btn.configure(
                text='시작점 고정 ✓' if fixed else '시작점 선택'
            )
            self.fiber_end_btn.configure(
                state='normal' if fixed else 'disabled'
            )
        except Exception:
            pass

    def _set_fiber_pick_mode(self,mode):
        if not self.fiber_network:
            return
        if mode=='start':
            # 시작점은 사용자가 이 버튼을 다시 누를 때만 변경한다.
            self.fiber_pick_mode='start'
            self.extract_hint.set('시작점 변경 모드 · 새 시작점을 1회 클릭하세요.')
        else:
            if self.route_start is None:
                self.fiber_pick_mode='start'
                self.extract_hint.set('먼저 시작점을 지정해주세요.')
            else:
                self.fiber_pick_mode='end'
                self.extract_hint.set('시작점 고정 · 끝점만 계속 클릭하면 경로가 누적됩니다.')
        self._update_fiber_pick_buttons()
        self.viewer.redraw()

    def _fiber_snapshot(self):
        return {
            'start':self.route_start,
            'routes':[dict(r) for r in self.fiber_routes],
            'mode':self.fiber_pick_mode,
        }

    def _push_fiber_history(self):
        self.fiber_history.append(self._fiber_snapshot())
        if len(self.fiber_history)>100:
            self.fiber_history=self.fiber_history[-100:]

    def _restore_fiber_snapshot(self,snapshot):
        self.route_start=snapshot.get('start')
        self.fiber_routes=[dict(r) for r in snapshot.get('routes') or []]
        self.fiber_pick_mode=snapshot.get('mode') or ('end' if self.route_start is not None else 'start')
        self._update_fiber_pick_buttons()
        self._refresh_fiber_selection()

    def _undo_fiber_selection(self):
        if not self.fiber_history:
            self.extract_hint.set('되돌릴 선택이 없습니다.')
            return
        self._restore_fiber_snapshot(self.fiber_history.pop())
        self.extract_hint.set(f'한 단계 취소 · 선택 경로 {len(self.fiber_routes)}개')

    def _clear_fiber_selections(self):
        if self.route_start is not None or self.fiber_routes:
            self._push_fiber_history()
        self.route_start=None
        self.fiber_routes=[]
        self.fiber_pick_mode='start'
        self._update_fiber_pick_buttons()
        self._refresh_fiber_selection()
        self.extract_hint.set('광케이블 선택 초기화 · 시작점 선택 버튼을 누르세요.')

    def _refresh_fiber_selection(self):
        if not self.fiber_network:
            return
        selected=set()
        for route in self.fiber_routes:
            selected.update(route.get('indices') or set())
        self.viewer.selected=selected
        self.viewer.route_start_marker=(
            self.fiber_network.positions.get(self.route_start)
            if self.route_start in self.fiber_network.positions else None
        )
        selected_ends=[]
        seen_ends=set()
        for r in self.fiber_routes:
            end_node=r.get('selection_end',r.get('end'))
            if end_node in seen_ends:
                continue
            seen_ends.add(end_node)
            if end_node in self.fiber_network.positions:
                selected_ends.append(self.fiber_network.positions.get(end_node))
        self.viewer.route_selected_markers=selected_ends
        if self.fiber_routes:
            all_indices=set().union(*(set(r.get('indices') or set()) for r in self.fiber_routes))
            self.export_scope={
                'kind':'fiber',
                'network':self.fiber_network,
                'routes':[{
                    'route':list(r.get('route') or []),
                    'nodes':list(r.get('nodes') or []),
                } for r in self.fiber_routes],
                'indices':all_indices,
            }
        else:
            self.export_scope=None
        self.viewer.redraw()

    def _run_catv_pole_excel(self):
        if not getattr(self,'can_extract',False): return
        if self.busy:return
        if not self.viewer.scene or not self.input_path:
            messagebox.showinfo('자가주 좌표 추출','먼저 DXF 도면을 열어주세요.');return
        rows,missing=catv_pole_rows(self.viewer.scene,self.input_path,int(self.epsg_var.get()))
        if not rows:
            messagebox.showinfo('자가주 좌표 추출',f'추출 가능한 Pole-CATV 자가주가 없습니다. 좌표 미확인 {len(missing)}개.');return
        handles={row[0] for row in rows}
        self._show_scope({'kind':'catv','indices':{e.index for e in self.viewer.scene.entities if e.handle in handles and e.entity_type=='INSERT'}})
        self.extract_hint.set(f'자가주 {len(rows)}개만 표시 / 지형·명칭 숨김')
        out=filedialog.asksaveasfilename(title='전체 자가주 좌표 Excel 추출',defaultextension='.xlsx',initialfile=Path(self.input_path).stem+'_자가주_좌표.xlsx',filetypes=[('Excel','*.xlsx')])
        if not out:return
        self.busy=True;self._show_progress();self._set_progress(10,'전체 자가주 삽입 좌표 추출')
        if missing:self._log(f'자가주 좌표 미확인 {len(missing)}개 제외: '+', '.join(missing))
        def worker():
            try:
                count=export_catv_poles(out,rows)
                self.q.put(('done',(ConversionStats(entities=count,rows=count),out)))
            except Exception as exc:self.q.put(('error',str(exc)))
        threading.Thread(target=worker,daemon=True).start()

    def _clear_network_state(self):
        self.viewer.hide_extraction_text=False
        self.export_scope=None;self.fiber_network=None;self.route_start=None;self.pipe_diameters={}
        self.fiber_routes=[];self.fiber_history=[];self.fiber_pick_mode=None
        self.viewer.entity_filter=None;self.viewer.network_click=None;self.viewer.route_nodes={};self.viewer.route_start_marker=None
        self.viewer.route_selected_markers=[];self.viewer.selected.clear()
        self.viewer.clear_fiber_fast_network()
        self._show_fiber_controls(False)
        self.extract_hint.set(
            '위·경도 좌표 추출 메뉴에서 작업을 선택하세요.'
            if bool(getattr(self,'can_extract',False))
            else '보기 전용 계정 · 추출/Excel 기능 없음'
        )

    def _reset_network_filter(self):
        self.viewer.hide_extraction_text=False
        self.export_scope=None;self.route_start=None
        self.fiber_routes=[];self.fiber_history=[];self.fiber_pick_mode=None
        self.viewer.entity_filter=None;self.viewer.network_click=None;self.viewer.route_nodes={};self.viewer.route_start_marker=None
        self.viewer.route_selected_markers=[]
        self.viewer.clear_fiber_fast_network()
        self._show_fiber_controls(False)
        self.viewer.set_mode('select');self.viewer.selected.clear();self.viewer.fit_view()
        self.extract_hint.set(
            '위·경도 좌표 추출 메뉴에서 작업을 선택하세요.'
            if bool(getattr(self,'can_extract',False))
            else '보기 전용 계정'
        )

    def _show_scope(self,scope):
        self.viewer.hide_extraction_text=True
        self.viewer.clear_fiber_fast_network()
        self.export_scope=scope
        display_indices=set(scope['indices'])
        self.viewer.entity_filter=display_indices
        layers={self.viewer.scene.entities[i].layer for i in display_indices}
        for iid in self.tree.get_children():
            name=self.layer_names.get(iid)
            if name in layers and not is_default_hidden(name):
                self.checked_layers.add(name);self.visible_layers.add(name);self._refresh_row(iid)
        self.viewer.visible_layers=set(self.visible_layers)
        self.viewer.network_click=None;self.viewer.route_nodes={};self.viewer.route_start_marker=None
        self.viewer.mode='select';self.viewer.selected.clear()
        # Fit and draw only extraction targets, once, without terrain or labels.
        self.viewer.fit_view()

    def _fiber_working_indices(self):
        """FiberNetwork에 실제 연결된 광케이블 + 연결 노드 기기만 표시."""
        if not self.fiber_network or not self.viewer.scene:
            return set()
        indices={edge.index for edge in self.fiber_network.edges.values()}
        adj=getattr(self.fiber_network,'adj',{}) or {}
        by_handle=getattr(self.fiber_network,'by_handle',{}) or {}
        positions=getattr(self.fiber_network,'positions',{}) or {}
        node_ids=set(adj.keys()) or set(positions.keys())
        for node in node_ids:
            entity=by_handle.get(node)
            if entity is not None:
                indices.add(entity.index)
        # NODELINK handle이 없는 CAD는 좌표로 연결된 함체만 보강한다.
        xy_points=[
            positions.get(node)
            for node in node_ids
            if str(node).startswith('XY')
        ]
        xy_points=[p for p in xy_points if p is not None]
        if xy_points:
            for entity in self.viewer.scene.entities:
                if not is_closure(entity):
                    continue
                pos=next(
                    (
                        tuple(data[:2])
                        for kind,data in (entity.primitives or [])
                        if kind in {'insert','point'}
                    ),
                    None,
                )
                if pos is None and entity.bbox:
                    x1,y1,x2,y2=entity.bbox
                    pos=((x1+x2)/2.0,(y1+y2)/2.0)
                if pos and any(math.dist(pos,p)<=0.1 for p in xy_points):
                    indices.add(entity.index)
        return indices

    def _activate_online_fiber_network(self,network,summary,selected_regions):
        self.busy=False
        self._set_progress(100,'온라인 광망 준비 완료')
        self.after(350,self._hide_progress)

        self.fiber_network=network
        self.export_scope=None
        self.route_start=None
        self.fiber_routes=[]
        self.fiber_history=[]
        self.fiber_pick_mode='start'

        self.viewer.hide_extraction_text=True
        self.viewer.entity_filter={edge.index for edge in network.edges.values()}
        self.viewer.selected.clear()
        self.viewer.network_click=self._fiber_route_click
        self.viewer.route_nodes=network.positions
        self.viewer.route_start_marker=None
        self.viewer.route_selected_markers=[]
        self.viewer.mode='fiber_route'
        self.viewer.set_fiber_fast_network(network)
        self._show_fiber_controls(True)
        self._update_fiber_pick_buttons()
        self.viewer.fit_fiber_fast()

        region_count=int((summary or {}).get('regions') or len(selected_regions or []))
        edge_count=len(network.edges)
        self.viewer.status_var.set(
            f'온라인 광추출 · 지역 {region_count}개 · 케이블 {edge_count:,}개'
        )
        self.extract_hint.set(
            '온라인 광망 준비 완료 · 시작점을 선택하면 고정되고 이후 끝점만 계속 선택합니다.'
        )

    def _begin_fiber_route(self):
        if not getattr(self,'can_extract',False): return
        if self.busy:return
        if not self.viewer.scene or not self.input_path:
            messagebox.showinfo('광케이블 추출','먼저 DXF 도면을 열어주세요.');return
        try:self.fiber_network=FiberNetwork(self.viewer.scene,self.input_path)
        except Exception as exc:messagebox.showerror('광케이블 추출',str(exc));return
        if not self.fiber_network.edges:
            messagebox.showinfo('광케이블 추출','지선을 제외한 연결 가능한 광케이블이 없습니다.');return
        self.viewer.hide_extraction_text=True
        self.export_scope=None
        self.route_start=None
        self.fiber_routes=[]
        self.fiber_history=[]
        self.fiber_pick_mode='start'
        # 선택 작업 성능을 위해 광케이블 + 실제 연결 기기만 남긴다.
        working_indices=self._fiber_working_indices()
        self.viewer.entity_filter=working_indices
        self.viewer.selected.clear()
        self.viewer.network_click=self._fiber_route_click
        self.viewer.route_nodes=self.fiber_network.positions
        self.viewer.route_start_marker=None
        self.viewer.route_selected_markers=[]
        self.viewer.mode='fiber_route'
        self.viewer.set_fiber_fast_network(self.fiber_network)
        self._show_fiber_controls(True)
        self._update_fiber_pick_buttons()
        self.viewer.redraw()
        self.viewer.status_var.set(
            f'광케이블 추출 · 광 연결 객체 {len(working_indices):,}개만 표시'
        )
        self.extract_hint.set('광케이블 추출 · 시작점 선택 버튼을 누른 뒤 국사/시작점을 클릭하세요.')

    def _nearest_fiber_node(self,event):
        hits=[]
        for node,point in self.fiber_network.positions.items():
            sx,sy=self.viewer.world_to_screen(*point)
            distance=math.hypot(event.x-sx,event.y-sy)
            if distance<=18:hits.append((distance,node))
        if not hits:
            self.extract_hint.set('표시된 선택점(함체/광케이블 끝점)을 클릭해주세요.')
            return None
        hits.sort()
        if len(hits)>1 and abs(hits[0][0]-hits[1][0])<1:
            self.extract_hint.set('선택점이 겹칩니다. 확대해서 구분되는 점을 선택해주세요.')
            return None
        return hits[0][1]

    def _fiber_new_segments(self, scope, selected_handles):
        """전체 경로에서 이미 선택된 케이블을 제거하고 새 연속구간만 반환."""
        route=list(scope.get('route') or [])
        nodes=list(scope.get('nodes') or [])
        if not route or len(nodes) != len(route) + 1:
            return []

        segments=[]
        start_i=None
        for i,handle in enumerate(route):
            is_new=handle not in selected_handles
            if is_new and start_i is None:
                start_i=i
            if start_i is not None and (not is_new or i == len(route)-1):
                end_i = i if (is_new and i == len(route)-1) else i-1
                seg_route=route[start_i:end_i+1]
                seg_nodes=nodes[start_i:end_i+2]
                seg_indices={self.fiber_network.edges[h].index for h in seg_route}
                # 분기 시작/끝을 포함한 연결 기기는 선택 표시 대상으로 유지한다.
                by_handle=getattr(self.fiber_network,'by_handle',{}) or {}
                for seg_node in seg_nodes:
                    ent=by_handle.get(seg_node)
                    if ent is not None and is_closure(ent):
                        seg_indices.add(ent.index)
                segments.append({
                    'route':seg_route,
                    'nodes':seg_nodes,
                    'indices':seg_indices,
                    'start':seg_nodes[0],
                    'end':seg_nodes[-1],
                })
                start_i=None
        return segments

    def _fiber_route_click(self,event):
        node=self._nearest_fiber_node(event)
        if node is None:
            return

        if self.fiber_pick_mode=='start':
            self._push_fiber_history()
            self.route_start=node
            self.viewer.route_start_marker=self.fiber_network.positions[node]
            # 한 번 선택한 시작점은 계속 고정한다.
            self.fiber_pick_mode='end'
            self._update_fiber_pick_buttons()
            self.extract_hint.set('시작점 고정 완료 · 이제 끝점만 계속 클릭하세요.')
            self._refresh_fiber_selection()
            return

        if self.route_start is None:
            self.fiber_pick_mode='start'
            self.extract_hint.set('먼저 시작점을 선택해주세요.')
            return

        # 같은 끝점을 다시 클릭하면 그 끝점 선택으로 추가된 모든 구간을 해제한다.
        existing=[
            i for i,r in enumerate(self.fiber_routes)
            if r.get('selection_end', r.get('end'))==node
        ]
        if existing:
            self._push_fiber_history()
            self.fiber_routes=[
                r for i,r in enumerate(self.fiber_routes)
                if i not in set(existing)
            ]
            self.fiber_pick_mode='end'
            self._update_fiber_pick_buttons()
            self._refresh_fiber_selection()
            endpoint_count=len({
                r.get('selection_end',r.get('end'))
                for r in self.fiber_routes
            })
            self.extract_hint.set(f'끝점 선택 해제 · 남은 끝점 {endpoint_count}개')
            return

        try:
            scope=self.fiber_network.unique_route(self.route_start,node)
        except ValueError as exc:
            messagebox.showwarning('광케이블 추출',str(exc));return

        selected_handles={
            h for route in self.fiber_routes
            for h in (route.get('route') or [])
        }
        full_route=list(scope.get('route') or [])
        overlap_count=len(selected_handles.intersection(full_route))
        new_segments=self._fiber_new_segments(scope,selected_handles)

        if not new_segments:
            self.extract_hint.set('선택한 끝점까지의 광케이블은 이미 모두 포함되어 있습니다.')
            self.viewer.status_var.set('추가할 새 광케이블 구간 없음')
            return

        self._push_fiber_history()
        for segment in new_segments:
            segment['selection_end']=node
            segment['fixed_start']=self.route_start
            self.fiber_routes.append(segment)

        self.fiber_pick_mode='end'
        self._update_fiber_pick_buttons()
        self._refresh_fiber_selection()

        cable_count=len({
            h for r in self.fiber_routes for h in (r.get('route') or [])
        })
        endpoint_count=len({
            r.get('selection_end',r.get('end'))
            for r in self.fiber_routes
        })
        added_count=sum(len(s['route']) for s in new_segments)
        branch_start=new_segments[0]['start']
        if overlap_count:
            self.extract_hint.set(
                f'공통구간 {overlap_count}개 자동 제외 · 분기점부터 새 구간 {added_count}개 추가 · '
                f'끝점 {endpoint_count}개'
            )
            self.viewer.status_var.set(
                f'중복 자동 제거 · 분기 시작 {branch_start} · 선택 광케이블 {cable_count}개'
            )
        else:
            self.extract_hint.set(
                f'새 구간 {added_count}개 추가 · 끝점 {endpoint_count}개 · '
                '시작점은 계속 고정되어 있습니다.'
            )
            self.viewer.status_var.set(f'선택 광케이블 {cable_count}개')

    def _begin_main_conduit(self):
        if not getattr(self,'can_extract',False): return
        if self.busy:return
        if not self.viewer.scene or not self.input_path:
            messagebox.showinfo('100mm 주관로 추출','먼저 DXF 도면을 열어주세요.');return
        self._reset_network_filter()
        records=conduit_records(self.viewer.scene,self.input_path,{})
        if not records:
            messagebox.showinfo('100mm 주관로 추출','읽을 수 있는 관로 객체가 없습니다.');return
        self._apply_main_conduit(records)

    def _apply_main_conduit(self,records):
        # 관경은 원본 CAD만 기준으로 사용하며 추출앱에서 보정하지 않는다.
        records=conduit_records(self.viewer.scene,self.input_path,{})
        main=[r for r in records if r['diameter']==100 and len(r['item']['points'])>=2]
        unknown=sum(r['diameter'] is None for r in records)
        drop=sum(r['diameter']==50 for r in records)
        other=sum(r['diameter'] not in (None,50,100) for r in records)
        missing=sum(len(r['item']['points'])<2 for r in records)
        if not main:
            messagebox.showwarning('100mm 주관로 추출',f'지름이 확인된 100mm 관로가 없습니다. 미확인 {unknown}개 / 50mm {drop}개는 추출하지 않습니다. 원본 CAD에서 관경 정보를 수정한 후 도면을 다시 읽어주세요.');return False
        self._show_scope({'kind':'pipe','records':records,'indices':{r['entity'].index for r in main}})
        self.extract_hint.set(f'100mm {len(main)}개 / 50mm {drop}·기타 {other}·미기재 {unknown} / 형상 미확인 {missing}')
        return True

    @staticmethod
    def _point_in_polygon(point, polygon):
        x,y=point;inside=False;j=len(polygon)-1
        for i in range(len(polygon)):
            xi,yi=polygon[i];xj,yj=polygon[j]
            if ((yi>y)!=(yj>y)) and (x < (xj-xi)*(y-yi)/((yj-yi) or 1e-15)+xi):
                inside=not inside
            j=i
        return inside

    @staticmethod
    def _segments_intersect(a,b,c,d):
        def orient(p,q,r):
            v=(q[1]-p[1])*(r[0]-q[0])-(q[0]-p[0])*(r[1]-q[1])
            if abs(v)<1e-9:return 0
            return 1 if v>0 else 2
        o1,o2,o3,o4=orient(a,b,c),orient(a,b,d),orient(c,d,a),orient(c,d,b)
        return o1!=o2 and o3!=o4

    def _primitive_hits_polygon(self,kind,data,polygon):
        if kind in {'insert','point'}:
            try:return self._point_in_polygon((float(data[0]),float(data[1])),polygon)
            except Exception:return False
        if kind=='circle':
            try:
                x,y,r=float(data[0]),float(data[1]),abs(float(data[2]))
                if self._point_in_polygon((x,y),polygon):return True
                for px,py in polygon:
                    if (px-x)*(px-x)+(py-y)*(py-y)<=r*r:return True
            except Exception:pass
            return False
        if kind in {'line','polyline','polygon'}:
            try:pts=[(float(x),float(y)) for x,y in data]
            except Exception:return False
            if any(self._point_in_polygon(p,polygon) for p in pts):return True
            edges=list(zip(polygon,polygon[1:]+polygon[:1]))
            segs=list(zip(pts,pts[1:]))
            if kind=='polygon' and len(pts)>=3:segs.append((pts[-1],pts[0]))
            return any(self._segments_intersect(a,b,c,d) for a,b in segs for c,d in edges)
        return False

    def _entity_hits_polygon(self,e,polygon):
        if not e.bbox:return False
        minx,miny,maxx,maxy=e.bbox
        pminx=min(p[0] for p in polygon);pmaxx=max(p[0] for p in polygon)
        pminy=min(p[1] for p in polygon);pmaxy=max(p[1] for p in polygon)
        if maxx<pminx or minx>pmaxx or maxy<pminy or miny>pmaxy:return False
        if any(self._primitive_hits_polygon(k,d,polygon) for k,d in e.primitives):return True
        center=((minx+maxx)/2,(miny+maxy)/2)
        return self._point_in_polygon(center,polygon)

    def _begin_custom_area_extract(self):
        if not getattr(self,'can_extract',False):return
        if self.busy:return
        if not self.viewer.scene or not self.input_path:
            messagebox.showinfo('사용자 지정 영역 추출','먼저 DXF 도면을 열어주세요.');return
        self._reset_network_filter()
        self.extract_hint.set('영역 꼭짓점을 순서대로 클릭하세요. Enter=폐합/완료, ESC=취소')
        def finished(points):
            polygon_indices={
                e.index for e in self.viewer.scene.entities
                if self._entity_hits_polygon(e,points)
            }
            if not polygon_indices:
                messagebox.showinfo(
                    '사용자 지정 영역 추출',
                    '지정 영역 안에 추출 가능한 객체가 없습니다.',
                )
                self._reset_network_filter()
                return

            # 영역 지정 후 실제 포함 레이어를 추가 옵션으로 선택한다.
            layers=sorted({
                self.viewer.scene.entities[i].layer
                for i in polygon_indices
            })

            win=tk.Toplevel(self)
            win.title('사용자 지정 영역 · 레이어 선택')
            win.geometry('620x560')
            win.minsize(540,460)
            win.transient(self)
            win.grab_set()

            root=ttk.Frame(win,padding=12)
            root.pack(fill='both',expand=True)

            ttk.Label(
                root,
                text='사용자 지정 영역 레이어 선택',
                font=('Malgun Gothic',14,'bold'),
            ).pack(anchor='w')
            ttk.Label(
                root,
                text=(
                    f'영역 내 객체 {len(polygon_indices):,}개 · '
                    '추출할 레이어만 선택하세요.'
                ),
            ).pack(anchor='w',pady=(2,10))

            box=ttk.LabelFrame(root,text='레이어',padding=8)
            box.pack(fill='both',expand=True)

            canvas=tk.Canvas(box,highlightthickness=0)
            scroll=ttk.Scrollbar(box,orient='vertical',command=canvas.yview)
            body=ttk.Frame(canvas)
            body_win=canvas.create_window((0,0),window=body,anchor='nw')
            body.bind(
                '<Configure>',
                lambda e:canvas.configure(scrollregion=canvas.bbox('all')),
            )
            canvas.bind(
                '<Configure>',
                lambda e:canvas.itemconfigure(body_win,width=e.width),
            )
            canvas.configure(yscrollcommand=scroll.set)
            canvas.pack(side='left',fill='both',expand=True)
            scroll.pack(side='right',fill='y')

            vars_by_layer={}
            for layer in layers:
                var=tk.BooleanVar(value=True)
                vars_by_layer[layer]=var
                count=sum(
                    1 for i in polygon_indices
                    if self.viewer.scene.entities[i].layer==layer
                )
                ttk.Checkbutton(
                    body,
                    text=f'{layer} · {count:,}개',
                    variable=var,
                ).pack(anchor='w',pady=2)

            def set_all(value):
                for var in vars_by_layer.values():
                    var.set(bool(value))

            def apply_layers():
                selected_layers={
                    layer
                    for layer,var in vars_by_layer.items()
                    if var.get()
                }
                if not selected_layers:
                    messagebox.showinfo(
                        '사용자 지정 영역 추출',
                        '레이어를 하나 이상 선택해주세요.',
                        parent=win,
                    )
                    return

                indices={
                    i for i in polygon_indices
                    if self.viewer.scene.entities[i].layer in selected_layers
                }
                scope={
                    'kind':'custom',
                    'indices':indices,
                    'polygon':list(points),
                    'layers':selected_layers,
                }
                win.destroy()
                self._show_scope(scope)
                self.extract_hint.set(
                    f'사용자 지정 영역 · 레이어 {len(selected_layers)}개 · '
                    f'객체 {len(indices):,}개 표시 / Excel 추출을 누르세요.'
                )
                self.status_var.set(
                    f'사용자 지정 영역 추출 대상 {len(indices):,}개'
                )

            footer=ttk.Frame(root)
            footer.pack(fill='x',pady=(10,0))
            ttk.Button(
                footer,
                text='전체 선택',
                command=lambda:set_all(True),
            ).pack(side='left')
            ttk.Button(
                footer,
                text='전체 해제',
                command=lambda:set_all(False),
            ).pack(side='left',padx=(6,0))
            ttk.Button(
                footer,
                text='적용',
                command=apply_layers,
            ).pack(side='right')
            ttk.Button(
                footer,
                text='취소',
                command=lambda:(win.destroy(),self._reset_network_filter()),
            ).pack(side='right',padx=(0,6))

        self.viewer.begin_area_select(finished)

    def _run_custom_area_excel(self):
        if self.busy:return
        scope=self.export_scope or {}
        indices=set(scope.get('indices') or [])
        if not indices:
            messagebox.showinfo('사용자 지정 영역 Excel','추출 대상이 없습니다.');return
        out=filedialog.asksaveasfilename(
            title='사용자 지정 영역 Excel 추출',
            defaultextension='.xlsx',
            initialfile=Path(self.input_path).stem+'_사용자지정영역.xlsx',
            filetypes=[('Excel','*.xlsx')],
        )
        if not out:return
        self.busy=True;self._show_progress();self._set_progress(10,'사용자 지정 영역 정보 정리')
        scene=self.viewer.scene;epsg=int(self.epsg_var.get())
        def worker():
            try:
                transform=Transformer.from_crs(f'EPSG:{epsg}','EPSG:4326',always_xy=True)
                headers=['순수 객체ID','레이어','객체종류','블록','SEQ','X','Y','경도','위도','문자','전주정보','블록속성','XDATA']
                rows=[]
                for idx in sorted(indices):
                    e=scene.entities[idx]
                    points=[]
                    for kind,data in e.primitives:
                        if kind in {'line','polyline','polygon'}:
                            points=[(float(x),float(y)) for x,y in data];break
                        if kind in {'insert','point'}:
                            points=[(float(data[0]),float(data[1]))];break
                    if not points and e.bbox:
                        x1,y1,x2,y2=e.bbox;points=[((x1+x2)/2,(y1+y2)/2)]
                    if not points:points=[(None,None)]
                    pole=json.dumps(e.pole_info or {},ensure_ascii=False)
                    attrs=json.dumps(e.attributes or {},ensure_ascii=False)
                    xdata=json.dumps(dict(e.xdata or []),ensure_ascii=False)
                    for seq,(x,y) in enumerate(points,1):
                        lon=lat=None
                        if x is not None and y is not None:
                            try:lon,lat=transform.transform(x,y)
                            except Exception:pass
                        rows.append([e.handle,e.layer,e.entity_type,e.block_name,seq,x,y,lon,lat,e.text,pole,attrs,xdata])
                wb=Workbook();ws=wb.active;ws.title='사용자지정영역'
                ws.append(headers)
                for cell in ws[1]:
                    cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='176D85')
                for row in rows:ws.append(row)
                ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
                wb.save(out)
                self.q.put(('done',(ConversionStats(entities=len(indices),rows=len(rows)),out)))
            except Exception as exc:self.q.put(('error',str(exc)))
        threading.Thread(target=worker,daemon=True).start()

    def _run_network_excel(self, submission=False):
        if not getattr(self,'can_extract',False): return
        if self.busy:return
        scope=self.export_scope
        if not scope:return
        if submission and scope['kind']!='fiber':
            messagebox.showinfo('제출용 엑셀','먼저 광케이블 경로를 추출해주세요.')
            return
        if scope['kind']=='custom':
            self._run_custom_area_excel();return
        if scope['kind']=='catv':
            self._run_catv_pole_excel();return
        if scope['kind']=='pipe':
            scope={**scope,'records':conduit_records(self.viewer.scene,self.input_path,self.pipe_diameters)}
            scope['indices']={r['entity'].index for r in scope['records'] if r['diameter']==100 and len(r['item']['points'])>=2}
        shown=self.viewer.entity_filter
        online_fiber=bool(scope.get('online') or getattr(scope.get('network'),'online',False))
        if online_fiber and scope['kind']=='fiber':
            indices=set(scope.get('indices') or [])
        else:
            indices=network_export_indices(self.viewer.scene,scope,self.viewer.visible_layers,shown,self.viewer.entity_items)
        if not indices:
            messagebox.showinfo('필요 항목 Excel 추출','화면에 표시된 추출 대상이 없습니다.');return
        label='광케이블' if scope['kind']=='fiber' else '100mm_주관로'
        base_name=Path(self.input_path).stem if self.input_path else ('온라인_광망' if online_fiber else '추출')
        export_title='광선로 제출용 Excel' if submission else f'{label} Excel 추출'
        export_name=base_name+('_광선로_제출.xlsx' if submission else '_'+label+'.xlsx')
        out=filedialog.asksaveasfilename(title=export_title,defaultextension='.xlsx',initialfile=export_name,filetypes=[('Excel','*.xlsx')])
        if not out:return
        self.busy=True;self._show_progress();self._set_progress(10,'전체 100mm 관로와 시작·끝 시설 위치 추출' if scope['kind']=='pipe' else '화면에 표시된 필요 항목만 Excel 추출')
        scene=self.viewer.scene;epsg=int(self.epsg_var.get())
        source_path=self.input_path or ('ONLINE' if online_fiber else '')
        def worker():
            try:
                exporter=export_submission_network if submission else export_network
                rows=exporter(out,scene,scope,indices,epsg,source_path=source_path)
                if scope['kind']=='fiber':
                    handles={
                        h for route in (scope.get('routes') or [])
                        for h in (route.get('route') or [])
                    }
                    count=sum(scope['network'].edges[h].index in indices for h in handles)
                else:
                    count=sum(r['diameter']==100 and r['entity'].index in indices for r in scope['records'])
                stats=ConversionStats(entities=count,rows=rows)
                self.q.put(('done',(stats,out)))
            except Exception as exc:self.q.put(('error',str(exc)))
        threading.Thread(target=worker,daemon=True).start()
