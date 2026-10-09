from pathlib import Path
from tempfile import TemporaryDirectory
import ezdxf
from openpyxl import load_workbook
from drawing_identity import COUNTIES, drawing_code, export_object_id
from converter import convert_selected_layers
from network_extract import FiberNetwork, export_network
from viewer import VisualEntity, Scene

for code in COUNTIES:
    assert drawing_code(f'CMB_GN_{code}_수정.dxf') == f'CMB_GN_{code}'
    assert export_object_id(f'CMB_GN_{code}.dxf', 'AB') == f'CMB_GN_{code}_AB'
assert drawing_code('cmb_gn_dy (2).dxf') == 'CMB_GN_DY'
assert export_object_id('CMB_GN_DY.dxf', 'AB/child/1') == 'CMB_GN_DY_AB/child/1'
assert export_object_id('CMB_GN_DY.dxf', 'CMB_GN_DY_AB') == 'CMB_GN_DY_AB'
assert export_object_id('CMB_GN_DY.dxf', '') == ''
with TemporaryDirectory() as temporary:
    root = Path(temporary)
    doc = ezdxf.new()
    entity = doc.modelspace().add_line((0, 0), (10, 0), dxfattribs={'layer': 'FIBER'})
    ids = []
    for code in ('DY', 'HP'):
        source = root / f'CMB_GN_{code}_수정.dxf'
        doc.saveas(source)
        output = root / f'{code}.xlsx'
        convert_selected_layers(source, output, ['FIBER'], 5174)
        workbook = load_workbook(output)
        ids.append(workbook['FIBER']['C2'].value)
        assert workbook['FIBER']['B1'].value == '순수 객체ID'
        assert workbook['FIBER']['C1'].value == '지역객체ID'
        assert workbook['FIBER']['B2'].value == entity.dxf.handle
        assert workbook['FIBER']['C2'].value == f'CMB_GN_{code}_{entity.dxf.handle}'
        assert workbook['FIBER']['B3'].value == workbook['FIBER']['B2'].value
    assert ids[0] != ids[1]
    source = root / 'CMB_GN_JS.dxf'
    source.write_text('0\nEOF\n')
    cable = VisualEntity(0, 'LINE', 'CN_F_Cable_FOC', 'AB', primitives=[('line', [(0, 0), (10, 0)])], bbox=(0, 0, 10, 0), xdata=[('EXMAP_NODELINK', ['1005:A', '1005:B'])])
    scene = Scene([cable], (0, 0, 10, 0), {})
    network = FiberNetwork(scene, source)
    scope = network.unique_route('A', 'B')
    export_network(root / 'fiber.xlsx', scene, scope, {0}, source_path=source)
    ws=load_workbook(root / 'fiber.xlsx')['광주간선']
    assert (ws['H1'].value,ws['I1'].value)==('순수 객체ID','지역객체ID')
    assert ws['H2'].value == 'AB' and ws['I2'].value == 'CMB_GN_JS_AB'
    assert ws['J1'].value=='시작경도' and isinstance(ws['J2'].value,float)
    assert ws['J2'].number_format=='0.0000000'
    pipe_scope = {'kind': 'pipe', 'indices': {0}, 'records': [{'entity': cable, 'diameter': 100, 'item': {'points': [(0, 0), (10, 0)]}}]}
    export_network(root / 'pipe.xlsx', scene, pipe_scope, {0}, source_path=source)
    ws=load_workbook(root / 'pipe.xlsx')['100mm_주관로']
    assert (ws['A1'].value,ws['B1'].value)==('순수 객체ID','지역객체ID')
    assert ws['A2'].value=='AB' and ws['B2'].value == 'CMB_GN_JS_AB'
    assert ws['C2'].value==100 and ws['I1'].value=='시작경도'
    assert isinstance(ws['I2'].value,float) and ws['I2'].number_format=='0.0000000'
    assert cable.handle == 'AB'  # Internal graph identifiers retain original handles.
print('IDENTITY TEST OK: 9 counties; file suffixes; general/fiber/pipe Excel; cross-county uniqueness; vertex rows retained')
