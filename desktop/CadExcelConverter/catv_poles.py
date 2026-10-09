import math
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pyproj import Transformer
from drawing_identity import export_object_id
from pipe_endpoints import facility_name

def is_catv_pole(entity):
    if entity.entity_type != 'INSERT':
        return False
    return any(re.search(r'(?:^|_)POLE_CATV(?:$|_)', str(value or '').upper().replace('-', '_'))
               for value in (entity.layer, entity.block_name))

def catv_pole_rows(scene, source_path, epsg=5174):
    transform = Transformer.from_crs(f'EPSG:{epsg}', 'EPSG:4326', always_xy=True)
    rows, missing = [], []
    for entity in scene.entities:
        if not is_catv_pole(entity):
            continue
        numbers = re.findall(r'[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?', str(entity.dxf_data.get('insert', '')))
        if len(numbers) < 2 or not all(math.isfinite(float(value)) for value in numbers[:2]):
            missing.append(entity.handle)
            continue
        x, y = map(float, numbers[:2])
        lon, lat = transform.transform(x, y)
        if not all(math.isfinite(value) for value in (lon, lat)):
            missing.append(entity.handle)
            continue
        rows.append([entity.handle, export_object_id(source_path, entity.handle),
                     facility_name(entity), x, y, lon, lat, entity.layer, entity.block_name])
    return rows, missing

def export_catv_poles(output_path, rows):
    if not rows:
        raise ValueError('좌표를 확인할 수 있는 자가주가 없습니다.')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '자가주_좌표'
    sheet.append(['순수 객체ID', '지역객체ID', '전주코드', 'CAD_X', 'CAD_Y', '경도', '위도', '레이어', '블록명'])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176D85')
    for row in sheet.iter_rows(min_row=2):
        for column in (4, 5):
            row[column-1].number_format = '0.000000'
        for column in (6, 7):
            row[column-1].number_format = '0.0000000'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(45, max(16, max(len(str(cell.value or '')) for cell in cells) + 2))
    workbook.save(output_path)
    return len(rows)
