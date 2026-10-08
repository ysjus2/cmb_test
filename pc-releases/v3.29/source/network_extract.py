from __future__ import annotations
import math
import re
from collections import deque
from dataclasses import dataclass
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pyproj import Transformer
from essenpoly_recovery import recover_essenpoly_polylines, recover_linker_polylines


def xvalues(entity, appid, code):
    rows = dict(entity.xdata or []).get(appid, [])
    prefix = str(code) + ':'
    return [str(v).split(':', 1)[1].strip() for v in rows if str(v).startswith(prefix)]


def entity_position(entity):
    value = entity.dxf_data.get('insert', '')
    nums = re.findall(r'[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?', value)
    if len(nums) >= 2:
        return float(nums[0]), float(nums[1])
    for kind, data in entity.primitives:
        if kind in {'insert', 'point'}:
            return tuple(data[:2])
    if entity.bbox:
        x1,y1,x2,y2 = entity.bbox
        return (x1+x2)/2, (y1+y2)/2
    return None


def is_fiber(entity):
    name = entity.layer.upper()
    return any(k in name for k in ('F_CABLE', 'FOC', 'FIBER', '광케이블'))


def is_closure(entity):
    return entity.entity_type == 'INSERT' and ('CLOSURE' in entity.layer.upper() or '함체' in entity.layer)


def is_background(entity):
    name = str(entity.layer or '').upper()
    if is_fiber(entity) or any(word in name for word in ('CABLE', 'CONDUIT', 'PIPELINE', '관로', '케이블', '전주', '함체')):
        return False
    if name.startswith('CN_'):
        return name in {'CN_M_USER_BUILDING', 'CN_M_USER_REGION'}
    return True


def cable_length(points):
    return sum(math.dist(a,b) for a,b in zip(points,points[1:]))


@dataclass
class FiberEdge:
    handle: str
    u: str
    v: str
    points: list
    index: int
    attributes: dict


class FiberNetwork:
    def __init__(self, scene, input_path):
        self.scene = scene
        self.edges = {}
        self.positions = {}
        self.adj = {}
        self.by_handle = {e.handle:e for e in scene.entities}
        recovered = {r['handle']:r for r in recover_essenpoly_polylines(input_path)}
        self.excluded_stubs = 0
        self.geometry_nodes = []
        def coordinate_node(point):
            for key, existing in self.geometry_nodes:
                if math.dist(point,existing) <= 0.1:
                    return key
            key = 'XY' + str(len(self.geometry_nodes))
            self.geometry_nodes.append((key,point))
            return key
        for e in scene.entities:
            if not is_fiber(e):continue
            item = recovered.get(e.handle)
            attrs = (item or {}).get('attributes',e.attributes)
            xdata = (item or {}).get('xdata',dict(e.xdata or []))
            if 'STUB' in (e.layer + ' ' + str(attrs.get('ESSEN_301',''))).upper() or '지선' in e.layer or 'EXMAP_ONULINK' in xdata:
                self.excluded_stubs += 1
                continue
            points = (item or {}).get('points')
            if not points:
                points = next((p for kind,p in e.primitives if kind in {'polyline','line'}),[])
            if len(points)<2:continue
            links = [v.split(':',1)[1].strip().upper() for v in xdata.get('EXMAP_NODELINK',[]) if str(v).startswith('1005:')]
            links = [v for v in links if v and v != '0']
            if len(links)==2:
                u,v=links
            else:
                u,v=coordinate_node(points[0]),coordinate_node(points[-1])
            if u==v:continue
            edge=FiberEdge(e.handle,u,v,points,e.index,attrs)
            self.edges[e.handle]=edge
            for node,point in ((u,points[0]),(v,points[-1])):
                target=self.by_handle.get(node)
                self.positions[node]=(entity_position(target) if target else None) or tuple(point)
                self.adj.setdefault(node,[]).append((v if node==u else u,e.handle))

    def unique_route(self,start,end):
        if start==end:raise ValueError('시작점과 끝점이 같습니다. 다른 함체/끝점을 선택해주세요.')
        def find(skip=None):
            parents={start:None};q=deque([start])
            while q:
                n=q.popleft()
                if n==end:
                    route=[];nodes=[end]
                    while parents[n] is not None:
                        previous,handle=parents[n];route.append(handle);nodes.append(previous);n=previous
                    return list(reversed(route)),list(reversed(nodes))
                for to,handle in self.adj.get(n,[]):
                    if handle!=skip and to not in parents:
                        parents[to]=(n,handle);q.append(to)
            return None
        found=find()
        if not found:raise ValueError('두 점 사이에 연결된 광케이블 경로가 없습니다. 지선(STUB)은 제외됩니다.')
        route,nodes=found
        if any(find(handle) is not None for handle in route):
            raise ValueError('두 점 사이에 여러 광경로가 있습니다. 두 점만으로 주간선을 확정할 수 없어 추출하지 않습니다. 더 짧은 구간의 함체를 선택해주세요.')
        closures={e.index for node in nodes for e in [self.by_handle.get(node)] if e and is_closure(e)}
        # Exact coordinate association is only used for drawings without node handles.
        for node in nodes:
            if not node.startswith('XY'):continue
            closures.update(e.index for e in self.scene.entities if is_closure(e) and entity_position(e) and math.dist(entity_position(e),self.positions[node])<=0.1)
        return {'kind':'fiber','route':route,'nodes':nodes,'network':self,'indices':{self.edges[h].index for h in route}|closures}


def conduit_records(scene,input_path,diameters=None):
    diameters=diameters or {}
    by_handle={e.handle:e for e in scene.entities}
    records=[]
    for item in recover_linker_polylines(input_path):
        if 'CONDUIT' not in item['layer'].upper():continue
        e=by_handle.get(item['handle'])
        if e is None:continue
        codes=item.get('xdata',{}).get('EXMAP_PIPELINE',[])
        code=next((v.split(':',1)[1].strip() for v in codes if v.startswith('1000:')),'')
        diameter=None
        # Never treat a cable length or an undocumented numeric XDATA field as diameter.
        for key,value in {**e.attributes,**item.get('attributes',{})}.items():
            if any(k in key.lower() for k in ('diameter','관경','지름','직경')):
                match=re.fullmatch(r'\s*(100|50)\s*(?:mm|㎜)?\s*',str(value),re.I)
                if match:diameter=int(match.group(1))
        if diameter is None:
            strings=[v.split(':',1)[1] for v in codes if v.startswith('1000:')]
            matches={int(m) for text in strings for m in re.findall(r'(?<!\d)(100|50)\s*(?:mm|㎜)',text,re.I)}
            if len(matches)==1:diameter=matches.pop()
        if diameter is None:diameter=diameters.get(item['handle'])
        records.append({'entity':e,'item':item,'code':code,'diameter':diameter})
    return records


def _node_name(network,node):
    e=network.by_handle.get(node)
    if e:
        values=xvalues(e,'EXMAP_FIBER',1000)
        if len(values)>1 and values[1]:return values[1]
        for key,value in e.attributes.items():
            if value and any(k in key.upper() for k in ('ID','NAME','번호')):return str(value)
    return node


def export_network(path,scene,scope,visible_indices,epsg=5174):
    visible_indices=set(visible_indices) & set(scope.get("indices", visible_indices))
    wb=Workbook();wb.remove(wb.active)
    rows=0
    geometry=[]
    def sheet(name,headers,data):
        nonlocal rows
        ws=wb.create_sheet(name);ws.append(headers)
        for row in data:ws.append(row);rows+=1
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='176D85')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(42,max(14,max(len(str(c.value or '')) for c in col)+2))
    transform=Transformer.from_crs(f'EPSG:{epsg}','EPSG:4326',always_xy=True)
    if scope['kind']=='fiber':
        network=scope['network'];data=[]
        for order,handle in enumerate(scope['route'],1):
            edge=network.edges[handle]
            if edge.index not in visible_indices:continue
            a,b=scope['nodes'][order-1:order+1]
            summary=str(edge.attributes.get('ESSEN_300','')).split('/')
            cable_id=edge.attributes.get('ESSEN_302') or summary[0] or handle
            cores=summary[1] if len(summary)>1 else ''
            declared=None
            if len(summary)>2:
                try:declared=float(summary[2])
                except ValueError:pass
            points = edge.points if edge.u == a else list(reversed(edge.points))
            start_point, end_point = points[0], points[-1]
            # Route endpoints are actual cable endpoints. Interior closures use
            # their own insertion positions instead of polyline bend vertices.
            start_closure = network.by_handle.get(a)
            end_closure = network.by_handle.get(b)
            if order > 1 and start_closure and is_closure(start_closure):
                start_point = entity_position(start_closure) or start_point
            if order < len(scope['route']) and end_closure and is_closure(end_closure):
                end_point = entity_position(end_closure) or end_point
            start_lon, start_lat = transform.transform(*start_point)
            end_lon, end_lat = transform.transform(*end_point)
            data.append([order,cable_id,cores,declared,round(cable_length(edge.points),3),_node_name(network,a),_node_name(network,b),handle,start_lon,start_lat,end_lon,end_lat])
        if not data:raise ValueError('화면에 표시된 주간선 광케이블이 없습니다.')
        sheet('광주간선',['순서','케이블번호','심수','도면기재길이_m','경로길이_m','시작함체','끝함체','객체ID','시작경도','시작위도','끝경도','끝위도'],data)
        ws=wb['광주간선']
        for row in ws.iter_rows(min_row=2,min_col=9,max_col=12):
            for cell in row:cell.number_format='0.0000000'
    else:
        from pipe_endpoints import PipeEndpointResolver
        resolver=PipeEndpointResolver(scene)
        data=[]
        for r in scope['records']:
            if r['diameter']!=100 or r['entity'].index not in visible_indices:continue
            pts=r['item']['points']
            start=resolver.resolve(pts[0]);end=resolver.resolve(pts[-1])
            start_lon,start_lat=transform.transform(*start['point'])
            end_lon,end_lat=transform.transform(*end['point'])
            data.append([r['entity'].handle,100,round(cable_length(pts),3),start['type'],start['name'],start_lon,start_lat,end['type'],end['name'],end_lon,end_lat])
        if not data:raise ValueError('화면에 표시된 100mm 주관로가 없습니다. 지름 미확인/50mm 관로는 추출하지 않습니다.')
        sheet('100mm_주관로',['객체ID','지름_mm','길이_m','시작시설','시작시설명_전주코드','시작경도','시작위도','끝시설','끝시설명_전주코드','끝경도','끝위도'],data)
        for row in wb['100mm_주관로'].iter_rows(min_row=2):
            for col in (6,7,10,11):row[col-1].number_format='0.0000000'
    wb.properties.description = f"CAD 좌표계 EPSG:{epsg}; 경위도 EPSG:4326"
    wb.save(path)
    return rows
