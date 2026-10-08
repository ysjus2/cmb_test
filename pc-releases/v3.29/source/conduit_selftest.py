from pathlib import Path
import tempfile

from openpyxl import load_workbook

from essenpoly_recovery import recover_linker_polylines
from converter import convert_selected_layers
from viewer import build_scene


def main():
    sample = """  0\nSECTION\n  2\nENTITIES\n  0\nASDKESSENLINKER\n  5\nABC1\n100\nAcDbEntity\n  8\nCN_L_Pole_Line_Conduit\n300\nCABLE-A\n101\nEmbedded Object\n100\nAcDbEntity\n  8\nCN_L_Pole_Line_Conduit\n 62\n5\n100\nAcDbPolyline\n 90\n        3\n 70\n     0\n 10\n200000.0\n 20\n500000.0\n 10\n200010.0\n 20\n500005.0\n 10\n200020.0\n 20\n500010.0\n1001\nEXMAP_NODELINK\n1000\nNODE-A\n1000\nNODE-B\n  0\nENDSEC\n  0\nEOF\n"""
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.dxf"
        xlsx = Path(td) / "sample.xlsx"
        path.write_text(sample, encoding="utf-8")

        rows = recover_linker_polylines(path)
        assert len(rows) == 1, rows
        row = rows[0]
        assert row["handle"] == "ABC1", row
        assert row["layer"] == "CN_L_Pole_Line_Conduit", row
        assert row["points"] == [
            (200000.0, 500000.0),
            (200010.0, 500005.0),
            (200020.0, 500010.0),
        ], row
        assert row["attributes"]["ESSEN_300"] == "CABLE-A", row
        assert row["color_aci"] == 5, row

        scene = build_scene(path)
        links = [e for e in scene.entities if e.entity_type == 'ASDKESSENLINKER']
        assert len(links) == 1, links
        assert links[0].layer == row['layer']
        assert links[0].primitives == [('polyline', row['points'])]
        assert links[0].color == '#0000ff', links[0].color
        assert not scene.unsupported.get('ASDKESSENLINKER')

        stats = convert_selected_layers(
            path,
            xlsx,
            ["CN_L_Pole_Line_Conduit"],
            source_epsg=5174,
        )
        assert stats.entities == 1, stats
        assert stats.rows == 3, stats

        wb = load_workbook(xlsx, data_only=True)
        ws = wb["CN_L_Pole_Line_Conduit"]
        assert ws.max_row == 4, ws.max_row
        assert ws["A2"].value == "ASDKESSENLINKER", ws["A2"].value
        assert ws["B2"].value == "ABC1", ws["B2"].value
        assert ws["D2"].value == 1, ws["D2"].value
        assert ws["E2"].value == 200000.0, ws["E2"].value
        assert ws["F2"].value == 500000.0, ws["F2"].value
        assert "CABLE-A" in str(ws["K2"].value), ws["K2"].value

        idx = wb["LAYER_INDEX"]
        assert idx["C2"].value == 1, idx["C2"].value
        assert idx["D2"].value == 3, idx["D2"].value

    print("ASDKESSENLINKER recovery + Excel export self-test OK")


if __name__ == "__main__":
    main()
