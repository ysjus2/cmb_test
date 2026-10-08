from pathlib import Path
import tempfile
import ezdxf
from openpyxl import load_workbook
from viewer import build_scene, _entity_primitives
from converter import scan_layers, convert_selected_layers


def main():
    with tempfile.TemporaryDirectory() as td:
        src = Path(td)/'all-layers.dxf'
        doc = ezdxf.new('R2018')
        for name in ('PARENT','LABEL','EMPTY','CHILD'):
            doc.layers.new(name)
        block = doc.blocks.new('SYMBOL')
        curve = block.add_lwpolyline([(0,0,1),(10,0,0)],format='xyb')
        hatch = block.add_hatch()
        hatch.paths.add_polyline_path([(0,0),(10,0),(10,10),(0,10)],is_closed=True)
        hatch.paths.add_polyline_path([(2,2),(2,8),(8,8),(8,2)],is_closed=True,flags=0)
        block.add_line((0,0),(1,1),dxfattribs={'layer':'CHILD'})
        ins = doc.modelspace().add_blockref('SYMBOL',(100,100),dxfattribs={'layer':'PARENT'})
        ins.add_attrib('ID','NODE-A',(101,102),dxfattribs={'layer':'LABEL'})
        doc.saveas(src)
        scene = build_scene(src)
        assert not scene.unsupported and not scene.diagnostics, (scene.unsupported,scene.diagnostics)
        assert {'PARENT','LABEL','CHILD'} <= {e.layer for e in scene.entities}
        assert 'EMPTY' in {l.name for l in scan_layers(src)}
        parent = next(e for e in scene.entities if e.entity_type=='INSERT')
        assert any(k=='polyline' and len(p)>2 and any(abs(y-100)>0.1 for x,y in p) for k,p in parent.primitives)
        triangles = [p for k,p in parent.primitives if k=='polygon']
        area = sum(abs((p[1][0]-p[0][0])*(p[2][1]-p[0][1])-(p[2][0]-p[0][0])*(p[1][1]-p[0][1]))/2 for p in triangles)
        assert abs(area-64)<1e-5, area
        assert any(e.entity_type=='ATTRIB' and e.text=='NODE-A' for e in scene.entities)
        out=Path(td)/'labels.xlsx'
        stats=convert_selected_layers(src,out,['LABEL'])
        assert stats.rows==1
        ws=load_workbook(out,data_only=True)['LABEL']
        assert ws['J2'].value=='NODE-A'
        assert 'ID=NODE-A' in ws['K2'].value
        doc.modelspace().add_mesh()
        doc.saveas(src)
        broken=build_scene(src)
        assert broken.unsupported.get('MESH')==1
        assert any(i['type']=='MESH' for i in broken.diagnostics)
    print('Full layer test OK: empty/child layers, attribute export, curved paths, hatch holes, unsupported diagnostics')

if __name__=='__main__':
    main()
