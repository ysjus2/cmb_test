from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from openpyxl import load_workbook
from viewer import VisualEntity, Scene
from network_extract import FiberNetwork, conduit_records, export_network


def fiber(i,h,u,v,a,b,layer='CN_F_Cable_FOC'):
    return VisualEntity(index=i,entity_type='LINE',layer=layer,handle=h,primitives=[('line',[a,b])],bbox=(min(a[0],b[0]),min(a[1],b[1]),max(a[0],b[0]),max(a[1],b[1])),xdata=[('EXMAP_NODELINK',['1005:'+u,'1005:'+v])])

def main():
    with TemporaryDirectory() as td:
        path=Path(td)/'network.dxf';path.write_text('0\nEOF\n')
        entities=[fiber(0,'AB','A','B',(0,0),(10,0)),fiber(1,'BC','B','C',(10,0),(20,0)),fiber(2,'BD','B','D',(10,0),(10,10)),fiber(3,'BS','B','S',(10,0),(10,-10),'CN_F_Cable_FOC-STUB')]
        for h,point in [('A',(0,0)),('B',(10,0)),('C',(20,0)),('D',(10,10))]:
            entities.append(VisualEntity(index=len(entities),entity_type='INSERT',layer='CN_F_Closure',handle=h,dxf_data={'insert':f'{point[0]}, {point[1]}, 0'},primitives=[('insert',(*point,'BOX'))],bbox=(*point,*point)))
        scene=Scene(entities,(0,-10,20,10),{})
        net=FiberNetwork(scene,path);scope=net.unique_route('A','C')
        assert scope['route']==['AB','BC']
        assert scope['indices']=={0,1,4,5,6}
        entities[5].dxf_data['insert']='10, 2, 0'
        out=Path(td)/'fiber.xlsx';export_network(out,scene,scope,scope['indices'])
        wb=load_workbook(out,data_only=True)
        assert wb['광주간선'].max_row==3 and '함체' not in wb.sheetnames
        assert wb.sheetnames==['광주간선']
        from pyproj import Transformer
        transform=Transformer.from_crs('EPSG:5174','EPSG:4326',always_xy=True)
        assert abs(wb['광주간선']['I2'].value-transform.transform(0,0)[0])<1e-9
        assert abs(wb['광주간선']['L3'].value-transform.transform(20,0)[1])<1e-9
        assert wb['광주간선']['K2'].value==wb['광주간선']['I3'].value
        assert wb['광주간선']['L2'].value==wb['광주간선']['J3'].value
        assert abs(wb['광주간선']['L2'].value-transform.transform(10,2)[1])<1e-9
        reverse=net.unique_route('C','A')
        export_network(out,scene,reverse,reverse['indices'])
        rev=load_workbook(out,data_only=True)['광주간선']
        assert abs(rev['I2'].value-transform.transform(20,0)[0])<1e-9
        assert abs(rev['L3'].value-transform.transform(0,0)[1])<1e-9
        assert wb['광주간선']['H2'].value=='AB' and wb['광주간선']['H3'].value=='BC'
        try:net.unique_route('A','NOT_CONNECTED')
        except ValueError:pass
        else:raise AssertionError('Disconnected route must be rejected')
        # A cycle makes two-point inference ambiguous and must not pick an arbitrary path.
        entities.append(fiber(len(entities),'AC','A','C',(0,0),(20,0)))
        try:FiberNetwork(scene,path).unique_route('A','C')
        except ValueError as exc:assert '여러' in str(exc)
        else:raise AssertionError('Ambiguous route must be rejected')
        raw=[];pipes=[]
        for h,code in [('AA1','100mm'),('AA2','50mm'),('AA3','PL0')]:
            raw.extend(['0','ASDKESSENLINKER','5',h,'8','CN_L_Pole_Line_Conduit','300','PL0|100.0','101','Embedded Object','100','AcDbEntity','8','CN_L_Pole_Line_Conduit','100','AcDbPolyline','90','2','10','0','20','0','10','10','20','0','1001','EXMAP_PIPELINE','1000',code])
            pipes.append(VisualEntity(index=len(pipes),entity_type='ASDKESSENLINKER',handle=h,layer='CN_L_Pole_Line_Conduit',primitives=[('line',[(0,0),(10,0)])]))
        path.write_text('\n'.join(raw+['0','EOF'])+'\n')
        pipe_scene=Scene(pipes,(0,0,10,0),{})
        records=conduit_records(pipe_scene,path)
        assert [r['diameter'] for r in records]==[100,50,None], records
        out=Path(td)/'pipes.xlsx';export_network(out,pipe_scene,{'kind':'pipe','records':records},{0,1,2})
        ws=load_workbook(out,data_only=True)['100mm_주관로']
        assert ws.max_row==2 and ws['A2'].value=='AA1'
        assert ws['B2'].value==100
        assert load_workbook(out,data_only=True).sheetnames==['100mm_주관로']
        assert ws['D2'].value=='미확인' and ws['H2'].value=='미확인'
        assert isinstance(ws['F2'].value,float)
        unknown=conduit_records(pipe_scene,path,{'AA3':100})
        assert unknown[2]['diameter']==100
    print('Network extraction test OK: off-route branches/STUB excluded; ambiguity/disconnection blocked; 50mm and unknown excluded')

if __name__=='__main__':main()
